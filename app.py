import os, io, copy, pickle, sqlite3, tempfile, uuid
from datetime import datetime
from zoneinfo import ZoneInfo
from PIL import Image
import pandas as pd
import streamlit as st

# Padrão SETTA: tabelas responsivas ao total de registros exibidos nos filtros.
def _setta_table_height(data, requested=None):
    """Altura baseada nas linhas reais (DataFrame ou pandas.Styler)."""
    if isinstance(requested, str):
        return requested
    frame = getattr(data, "data", data)
    try:
        rows = len(frame)
    except (TypeError, ValueError, AttributeError):
        return requested if isinstance(requested, int) and requested >= 120 else None
    ceiling = max(120, requested) if isinstance(requested, int) and requested > 0 else 600
    return int(min(ceiling, max(120, 42 + 35 * (min(max(0, rows), 100) + 1))))


def _setta_dataframe(data, *args, **kwargs):
    height = _setta_table_height(data, kwargs.get("height"))
    if height is None:
        kwargs.pop("height", None)
    else:
        kwargs["height"] = height
    return st.dataframe(data, *args, **kwargs)


def _setta_data_editor(data, *args, **kwargs):
    if kwargs.get("num_rows") != "dynamic":
        height = _setta_table_height(data, kwargs.get("height"))
        if height is None:
            kwargs.pop("height", None)
        else:
            kwargs["height"] = height
    return st.data_editor(data, *args, **kwargs)


import firebase_admin
from firebase_admin import credentials, firestore
import central_inventory_data as central_data
import setta_shell
import setta_auth

try:
 _GLOBAL_VISUAL_CONFIG=central_data.load_visual_config()
except Exception:
 _GLOBAL_VISUAL_CONFIG={}

SETTA_UI_CONFIG=setta_shell.build_ui_config(
 _GLOBAL_VISUAL_CONFIG.get('ui_config') or {}
)

def _global_browser_icon():
 try:
  raw=central_data.favicon_bytes(_GLOBAL_VISUAL_CONFIG)
  if raw:
   image=Image.open(io.BytesIO(raw))
   image.load()
   return image
 except Exception:
  pass
 return '📡'

st.set_page_config(
 page_title='GESTÃO DE ESTOQUE | SETTA',
 page_icon=_global_browser_icon(),
 layout='wide',
 initial_sidebar_state='expanded',
)

def _setta_sidebar_is_open():
 return bool(st.session_state.get('_setta_sidebar_open',False))

def _setta_toggle_sidebar():
 st.session_state['_setta_sidebar_open']=not _setta_sidebar_is_open()

def _setta_close_sidebar():
 st.session_state['_setta_sidebar_open']=False

TZ=ZoneInfo('America/Sao_Paulo')

def now_local():
 return datetime.now(TZ)

# O shell é emitido antes de qualquer leitura remota operacional.
setta_shell.render_shell(st,SETTA_UI_CONFIG,sidebar_open=_setta_sidebar_is_open())

BUILD_DIAGNOSTICO = 'inventory-20261008-gestor-ui-delete'
INVENTORY_DATA_EPOCH='20261007-RESET-01'
DATA=os.path.join(tempfile.gettempdir(),'inventario_operacional.sqlite3')

ESTOQUE_ENDERECOS_NAO_DISPONIVEIS = {
 'ALMOX. N.C.I',
 'ALMOX. N.C.R',
 'ASTEC',
 'G9-PROD',
 'G9-SEPA',
 'PROD. BARRAS',
 'PROD. CHAPAS',
 'PROD. PROC',
 'PROD. SUCATA',
 'PROD. TCTP',
 'QUALIDADE',
 'SETTA FIOS',
}

DEFAULT={
 'blind_default':False,
 'new_inventory_text':'NOVO INVENTÁRIO',
}

DIVERGENCE_REASONS=[
 'ERRO DE CONTAGEM',
 'MOVIMENTAÇÃO NÃO REGISTRADA',
 'ENTRADA NÃO REGISTRADA',
 'SAÍDA/CONSUMO NÃO REGISTRADO',
 'RETORNO DE PRODUÇÃO',
 'MATERIAL EM ENDEREÇO INCORRETO',
 'AVARIA/PERDA',
 'DIVERGÊNCIA DE SISTEMA',
 'OUTRO',
]
DIVERGENCE_TREATMENTS=['AJUSTAR ESTOQUE','ENCERRAR SEM AJUSTE']


def dbconn():
 c=sqlite3.connect(DATA); c.execute('CREATE TABLE IF NOT EXISTS state(k TEXT PRIMARY KEY,v BLOB)'); c.commit(); return c
def save(k,v):
 c=dbconn(); c.execute('INSERT OR REPLACE INTO state(k,v) VALUES(?,?)',(k,sqlite3.Binary(pickle.dumps(v)))); c.commit(); c.close()
def load(k,d=None):
 c=dbconn(); r=c.execute('SELECT v FROM state WHERE k=?',(k,)).fetchone(); c.close()
 if not r:return d
 try:return pickle.loads(r[0])
 except:return d



# Firebase Admin / Firestore: contingência somente leitura.
def firebase_db():
    try:
        if not firebase_admin._apps:
            sa = dict(st.secrets.get('firebase_admin', {}))
            if not sa.get('project_id') or not sa.get('private_key') or not sa.get('client_email'):
                return None
            firebase_admin.initialize_app(credentials.Certificate(sa))
        return firestore.client()
    except Exception:
        return None

def _fs_load_df(db, name):
    rows = [x.to_dict() for x in db.collection(name).stream()]
    if not rows:
        return None
    rows.sort(key=lambda x: x.get('_ordem', 0))
    for r in rows:
        r.pop('_ordem', None)
    return pd.DataFrame(rows)

def _fs_load_eligible(db):
    snap = db.collection('estoque_config').document('enderecos').get()
    if not snap.exists:
        return None
    return list((snap.to_dict() or {}).get('enderecos') or [])

def firestore_load_db():
    db = firebase_db()
    if db is None:
        return None
    try:
        return _fs_load_df(db, 'estoque_produtos')
    except Exception:
        return None

def firestore_load_pos():
    db = firebase_db()
    if db is None:
        return None
    try:
        return _fs_load_df(db, 'estoque_posicoes')
    except Exception:
        return None

def firestore_load_eligible():
    db = firebase_db()
    if db is None:
        return None
    try:
        return _fs_load_eligible(db)
    except Exception:
        return None


_remote_operational_state={}
_force_operational_reload=(
 st.session_state.get('_inventory_data_epoch')!=INVENTORY_DATA_EPOCH
)
if _force_operational_reload or any(
 key not in st.session_state
 for key in ('cfg','inventories','cycles','reports')
):
 try:
  _remote_operational_state=central_data.operational_state(
   ['cfg','inventories','cycles','reports']
  )
  st.session_state.pop('_operational_persistence_error',None)
 except Exception as exc:
  _remote_operational_state={}
  st.session_state['_operational_persistence_error']=str(exc)

def _initial_operational_value(key,default):
 if key in _remote_operational_state:
  value=_remote_operational_state.get(key)
  return value if value is not None else default
 return load(key,default)

if 'cfg' not in st.session_state:
 _saved_cfg=_initial_operational_value('cfg',{}) or {}
 st.session_state.cfg={**DEFAULT,**(_saved_cfg if isinstance(_saved_cfg,dict) else {})}

# A base de estoque é carregada da Central SETTA depois do shell.
# Firestore/SQLite ficam apenas como contingência se a Central falhar.
if 'db' not in st.session_state: st.session_state.db=None
if 'pos' not in st.session_state: st.session_state.pos=None
if 'eligible' not in st.session_state: st.session_state.eligible=[]

if _force_operational_reload:
 _saved_inventories=_remote_operational_state.get('inventories',{}) or {}
 _saved_cycles=_remote_operational_state.get('cycles',{}) or {}
 _saved_reports=_remote_operational_state.get('reports',{}) or {}
 st.session_state.inventories=_saved_inventories if isinstance(_saved_inventories,dict) else {}
 st.session_state.cycles=_saved_cycles if isinstance(_saved_cycles,dict) else {}
 st.session_state.reports=_saved_reports if isinstance(_saved_reports,dict) else {}
 st.session_state.selected=None
 st.session_state.new_inv=False
 st.session_state['_inventory_data_epoch']=INVENTORY_DATA_EPOCH
 save('inventories',st.session_state.inventories)
 save('cycles',st.session_state.cycles)
 save('reports',st.session_state.reports)
else:
 if 'inventories' not in st.session_state:
  _saved_inventories=_initial_operational_value('inventories',{}) or {}
  st.session_state.inventories=_saved_inventories if isinstance(_saved_inventories,dict) else {}
 if 'cycles' not in st.session_state:
  _saved_cycles=_initial_operational_value('cycles',{}) or {}
  st.session_state.cycles=_saved_cycles if isinstance(_saved_cycles,dict) else {}
 if 'reports' not in st.session_state:
  _saved_reports=_initial_operational_value('reports',{}) or {}
  st.session_state.reports=_saved_reports if isinstance(_saved_reports,dict) else {}

if 'section' not in st.session_state: st.session_state.section='Dashboard'
if 'selected' not in st.session_state: st.session_state.selected=None
if 'new_inv' not in st.session_state: st.session_state.new_inv=False

cfg=st.session_state.cfg
config=cfg

# Migração transparente do estado local antigo para a persistência atômica.
# Executa uma vez por sessão e apenas quando a chave ainda não existe remotamente.
if not st.session_state.get('_operational_state_migration_checked'):
 try:
  if 'inventories' not in _remote_operational_state and st.session_state.inventories:
   for _doc,_inv in st.session_state.inventories.items():
    central_data.save_inventory_document(str(_doc),_inv)

  if 'cycles' not in _remote_operational_state and st.session_state.cycles:
   central_data.merge_inventory_cycles(st.session_state.cycles)

  if 'reports' not in _remote_operational_state and st.session_state.reports:
   for _rid,_rep in st.session_state.reports.items():
    central_data.save_inventory_report(str(_rid),_rep)

  st.session_state.pop('_operational_persistence_error',None)
 except Exception as exc:
  st.session_state['_operational_persistence_error']=str(exc)
 st.session_state['_operational_state_migration_checked']=True

# Autenticação central SETTA. A política global do OperaHub decide se o login é obrigatório.
@st.cache_data(ttl=60,show_spinner=False,max_entries=2)
def _setta_auth_bootstrap_cached(anon_key):
 return setta_auth.bootstrap(anon_key)

def _auth_user():
 user=st.session_state.get('_setta_auth_user')
 return user if isinstance(user,dict) else None

def _session_operator():
 user=_auth_user()
 if user:
  return str(user.get('full_name') or user.get('username') or 'USUÁRIO SETTA').strip()
 return 'OPERADOR NÃO IDENTIFICADO'

def _session_profile():
 user=_auth_user()
 role=str((user or {}).get('role') or '').strip().lower()
 return ('Produção' if role=='producao' else 'Almoxarifado' if role=='almoxarifado' else
         'Gestor' if role in {'admin','gestor'} else 'Operador')

def _session_auth_token():
 user=_auth_user()
 return str((user or {}).get('inventory_token') or '').strip()

def _authenticate_user(login,password):
 user=setta_auth.authenticate(central_data.supabase_key(),login,password)
 if not user:
  return False
 st.session_state['_setta_auth_user']=user
 st.session_state.pop('_setta_auth_password',None)
 return True

def _setta_login_required():
 key=central_data.supabase_key()
 if not key:
  st.session_state['_setta_auth_policy_error']='SUPABASE_ANON_KEY não configurada.'
  return True
 try:
  payload=_setta_auth_bootstrap_cached(key)
  settings=payload.get('settings') if isinstance(payload,dict) else {}
  required=bool((settings or {}).get('login_required',False))
  st.session_state['_setta_auth_required']=required
  st.session_state.pop('_setta_auth_policy_error',None)
  return required
 except Exception as exc:
  st.session_state['_setta_auth_policy_error']=str(exc)
  return True

def _render_setta_auth_gate():
 if not _setta_login_required() or _auth_user():
  return True
 section_band('ACESSO SETTA','AUTENTICAÇÃO NECESSÁRIA','UTILIZE O MESMO USUÁRIO E SENHA DO OPERAHUB.')
 if st.session_state.get('_setta_auth_policy_error'):
  st.warning('Não foi possível confirmar a política central de acesso. Por segurança, o aplicativo permanece bloqueado até a autenticação.')
 with st.form('setta_auth_login_form',clear_on_submit=False):
  login=st.text_input('USUÁRIO OU E-MAIL',key='_setta_auth_login')
  password=st.text_input('SENHA',type='password',key='_setta_auth_password')
  submitted=st.form_submit_button('ENTRAR',type='primary',use_container_width=True)
 if submitted:
  try:
   user_ok=_authenticate_user(login,password)
  except Exception as exc:
   st.error(f'Não foi possível validar o acesso: {exc}')
  else:
   if user_ok:
    st.rerun()
   else:
    st.error('Usuário ou senha inválidos.')
 st.stop()
def persist_cfg():
 save('cfg',cfg)
 try:
  central_data.save_operational_state('cfg',cfg,_session_auth_token())
  st.session_state.pop('_operational_persistence_error',None)
  return True
 except Exception as exc:
  message=str(exc)
  if 'SESSION_INVALID_OR_EXPIRED' in message or 'AUTH_REQUIRED' in message:
   st.session_state.pop('_setta_auth_user',None)
  st.session_state['_operational_persistence_error']=message
  return False

def persist_report(report):
 rid=str(report.get('id') or '').strip()
 if not rid:return False
 st.session_state.reports[rid]=report
 save('reports',st.session_state.reports)
 try:
  central_data.save_inventory_report(rid,report)
  st.session_state.pop('_operational_persistence_error',None)
  return True
 except Exception as exc:
  st.session_state['_operational_persistence_error']=str(exc)
  return False

def excel_bytes(df, sheet_name):
 out=io.BytesIO()
 with pd.ExcelWriter(out,engine='openpyxl') as writer:
  df.to_excel(writer,index=False,sheet_name=sheet_name)
  ws=writer.book[sheet_name]
  ws.freeze_panes='A2'
  ws.auto_filter.ref=ws.dimensions
  from openpyxl.styles import Font, PatternFill, Alignment
  for cell in ws[1]:
   cell.font=Font(bold=True,color='11130F')
   cell.fill=PatternFill('solid',fgColor='FFD63B')
   cell.alignment=Alignment(horizontal='center',vertical='center')
  ws.row_dimensions[1].height=24
  for col in ws.columns:
   letter=col[0].column_letter
   max_len=max(len(str(c.value)) if c.value is not None else 0 for c in col[:80])
   ws.column_dimensions[letter].width=min(max(max_len+2,10),60)
 return out.getvalue()

def signed_brl(v):
 x=float(v)
 if abs(x)<1e-12:return 'R$ 0,00'
 return ('+' if x>0 else '-')+brl(abs(x))

def logo_uri():
 try:
  return central_data.logo_data_uri(_GLOBAL_VISUAL_CONFIG) or None
 except Exception:
  return None

def section_band(kicker,title,note=''):
 note_html=f'<div class="section-band-note">{note}</div>' if str(note or '').strip() else ''
 st.markdown(f'<div class="section-band"><div class="section-band-kicker">{kicker}</div><div class="section-band-title">{title}</div>{note_html}</div>',unsafe_allow_html=True)

def topic_divider():
 st.markdown('<div class="topic-divider"></div>',unsafe_allow_html=True)

def setta_kpi(
 col,
 label,
 value,
 delta='',
 accent='#111827',
 soft='#f3f4f6',
 extra='',
):
 # Mantém compatibilidade com chamadas antigas que ainda enviem um
 # sétimo argumento visual. O conteúdo extra é opcional e não interfere
 # no layout padrão do card.
 extra_html=(
  f'<div class="kpi-extra">{extra}</div>'
  if str(extra or '').strip()
  else ''
 )
 col.markdown(
  f'<div class="kpi-card" style="--accent:{accent};--accent-soft:{soft}">'
  f'<div class="kpi-header"><span class="kpi-dot"></span><span class="kpi-label">{label}</span></div>'
  f'<div class="kpi-value">{value}</div><div class="kpi-delta">{delta}</div>'
  f'{extra_html}</div>',
  unsafe_allow_html=True
 )

def dashboard_kpi_grid(items,valor_apto,qtd_cnt,qtd_div,acc_itens,acc_pos):
 cards=[
  ('blue','ITENS DIFERENTES COM SALDO',f'{items:,}'.replace(',','.'),'BASE APTA','ESTOQUE'),
  ('cyan','VALOR TOTAL APTO A CONTABILIZAR',brl(valor_apto),'VALOR DO ESTOQUE','FINANCEIRO'),
  ('green','POSIÇÕES CONTABILIZADAS',f'{qtd_cnt:,}'.replace(',','.'),'CONTAGENS REGISTRADAS','CONTAGEM'),
  ('red','POSIÇÕES DIVERGENTES',f'{qtd_div:,}'.replace(',','.'),'EXIGEM TRATATIVA','DIVERGÊNCIA'),
  ('violet','ACURÁCIA · ITENS COM SALDO',f'{acc_itens:.2f}%','ÍNDICE GERAL','ACURÁCIA'),
  ('amber','ACURÁCIA · POSIÇÕES CONTABILIZADAS',f'{acc_pos:.2f}%','ÍNDICE CONTABILIZADO','ACURÁCIA'),
 ]
 html=['<div class="dashboard-kpi-grid">']
 for tone,label,value,caption,eyebrow in cards:
  progress=''
  if tone in ('violet','amber'):
   pct=max(0.0,min(100.0,float(acc_itens if tone=='violet' else acc_pos)))
   progress=f'<div class="dashboard-kpi-progress"><span style="width:{pct:.2f}%"></span></div>'
  html.append(
   f'<div class="dashboard-kpi dashboard-kpi--{tone}">'
   f'<div class="dashboard-kpi-top"><span class="dashboard-kpi-eyebrow">{eyebrow}</span><span class="dashboard-kpi-status-dot"></span></div>'
   f'<div class="dashboard-kpi-label">{label}</div>'
   f'<div class="dashboard-kpi-value">{value}</div>'
   f'{progress}'
   f'<div class="dashboard-kpi-caption">{caption}</div>'
   f'</div>'
  )
 html.append('</div>')
 st.markdown(''.join(html),unsafe_allow_html=True)

 return None

st.markdown(
    """
<style id="inventory-component-style">
div[data-testid="stElementContainer"]:has(#inventory-component-style){display:none!important;height:0!important;min-height:0!important;margin:0!important;padding:0!important}
:root{--p:#111827}
.section-title{margin:0 0 1rem!important;color:#0f172a!important;font-size:1.28rem!important;font-weight:900!important;letter-spacing:-.02em;text-transform:uppercase}
.section-band{margin:0 0 .95rem!important;padding:.82rem 1rem!important;background:#fff!important;border:1px solid #e5e8ee!important;border-left:5px solid #111827!important;border-radius:12px!important;box-shadow:0 3px 12px rgba(15,23,42,.035)!important}
.section-band-kicker{font-size:.66rem!important;font-weight:900!important;letter-spacing:.085em!important;text-transform:uppercase!important;color:#ef4444!important;margin-bottom:.18rem!important}
.section-band-title{font-size:1.08rem!important;font-weight:900!important;color:#111827!important;letter-spacing:-.015em!important;line-height:1.2!important;text-transform:uppercase!important}
.section-band-note{margin-top:.22rem!important;color:#667085!important;font-size:.75rem!important;line-height:1.35!important}
.topic-divider{height:1px!important;background:#cbd5e1!important;margin:1.55rem 0 1.05rem!important;width:100%!important}
.api-status-head{display:flex;justify-content:space-between;align-items:center;gap:1rem;flex-wrap:wrap;margin-bottom:.8rem}
.api-status-name{font-size:.92rem;font-weight:900;color:#111827;text-transform:uppercase}
.api-status-filter{margin-top:.18rem;font-size:.7rem;color:#64748b;font-weight:700;text-transform:uppercase;letter-spacing:.025em}
.api-status-badge{display:inline-flex;padding:.28rem .52rem;border-radius:999px;background:#dcfce7;color:#166534;font-size:.68rem;font-weight:900;letter-spacing:.035em}
.api-grid{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:.7rem;margin:.25rem 0 .8rem}
.api-stat{background:#f8fafc;border:1px solid #e5e7eb;border-radius:10px;padding:.7rem .78rem}
.api-stat-label{font-size:.61rem;font-weight:900;letter-spacing:.055em;text-transform:uppercase;color:#64748b;margin-bottom:.28rem}
.api-stat-value{font-size:.92rem;font-weight:900;color:#111827;line-height:1.15;overflow-wrap:anywhere}
.setta-empty-state{border:1px dashed #cbd5e1;border-radius:12px;background:#f8fafc;padding:.85rem 1rem;color:#64748b;font-size:.76rem;font-weight:800;letter-spacing:.02em;text-transform:uppercase;margin:.1rem 0 .5rem}
.intro,.panel{background:#fff;border:1px solid #e5e8ee;border-radius:14px}
.intro{padding:.9rem 1rem;color:#555c66;margin-bottom:1rem;box-shadow:0 3px 12px rgba(15,23,42,.035)}
.kpi-card{position:relative;min-height:116px;padding:16px 18px 15px;border:1px solid #e2e8f0;border-radius:14px;background:#fff;box-shadow:0 4px 16px rgba(15,23,42,.055);overflow:hidden;transition:transform .12s ease,box-shadow .12s ease}
.kpi-card::before{content:"";position:absolute;left:0;top:0;bottom:0;width:5px;background:var(--accent)}
.kpi-card.selected{outline:2px solid var(--accent);outline-offset:1px}
.kpi-header{display:flex;align-items:center;gap:8px;margin-bottom:11px}
.kpi-dot{width:9px;height:9px;border-radius:999px;background:var(--accent);box-shadow:0 0 0 4px var(--accent-soft);flex:0 0 auto}
.kpi-label{color:#475569;font-size:.72rem!important;font-weight:900!important;line-height:1.15;text-transform:uppercase;letter-spacing:.025em}
.kpi-value{color:#0f172a;font-size:2rem;font-weight:800;line-height:1;letter-spacing:-.035em}
.kpi-delta{margin-top:8px;color:#64748b;font-size:.66rem!important;font-weight:700!important;text-transform:uppercase;letter-spacing:.02em}
[data-testid="stDataFrame"],[data-testid="stDataEditor"]{border:1px solid #d9dee7;border-radius:12px;overflow:hidden;box-shadow:0 5px 18px rgba(15,23,42,.06);background:#fff}
[data-testid="stDataFrame"] [role="columnheader"],[data-testid="stDataEditor"] [role="columnheader"]{font-weight:800!important;text-transform:uppercase!important;letter-spacing:.025em!important;background:#f8fafc!important}
[data-testid="stDataFrame"] [role="gridcell"],[data-testid="stDataEditor"] [role="gridcell"]{border-color:#eef1f5!important}
[data-testid="stVerticalBlockBorderWrapper"]{border-color:#e5e8ee!important;border-radius:14px!important;background:#fff!important;box-shadow:0 3px 12px rgba(15,23,42,.035)!important}
[data-testid="stTabs"] button{font-weight:800!important;text-transform:uppercase!important;letter-spacing:.015em!important}
div[data-testid="stMarkdownContainer"] h1,div[data-testid="stMarkdownContainer"] h2,div[data-testid="stMarkdownContainer"] h3,div[data-testid="stMarkdownContainer"] h4{text-transform:uppercase}
[data-testid="stAlert"]{border-radius:12px!important;box-shadow:0 3px 12px rgba(15,23,42,.035)!important}
button[kind="primary"],button[data-testid="stBaseButton-primary"]{background:#111827!important;border-color:#111827!important;color:#fff!important}
.footer{text-align:center;color:#9298a1;font-size:.72rem;padding-top:1.2rem}
.dashboard-kpi-grid{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:.85rem;margin:.15rem 0 .25rem}
.dashboard-kpi{--kpi-accent:#2563eb;--kpi-soft:#eff6ff;position:relative;min-width:0;min-height:142px;padding:1rem 1.05rem .9rem;background:#fff;border:1px solid #e2e8f0;border-radius:14px;box-shadow:0 5px 18px rgba(15,23,42,.055);overflow:hidden;display:flex;flex-direction:column}
.dashboard-kpi::before{content:"";position:absolute;left:0;top:0;bottom:0;width:4px;background:var(--kpi-accent)}
.dashboard-kpi--blue{--kpi-accent:#2563eb;--kpi-soft:#dbeafe}.dashboard-kpi--cyan{--kpi-accent:#0891b2;--kpi-soft:#cffafe}.dashboard-kpi--green{--kpi-accent:#16a34a;--kpi-soft:#dcfce7}.dashboard-kpi--red{--kpi-accent:#ef4444;--kpi-soft:#fee2e2}.dashboard-kpi--violet{--kpi-accent:#7c3aed;--kpi-soft:#ede9fe}.dashboard-kpi--amber{--kpi-accent:#d97706;--kpi-soft:#ffedd5}
.dashboard-kpi-top{display:flex;align-items:center;justify-content:space-between;gap:.7rem;margin-bottom:.48rem}
.dashboard-kpi-eyebrow{color:#64748b;font-size:.59rem;font-weight:900;letter-spacing:.085em;text-transform:uppercase}
.dashboard-kpi-status-dot{width:9px;height:9px;border-radius:999px;background:var(--kpi-accent);box-shadow:0 0 0 4px var(--kpi-soft);flex:0 0 auto}
.dashboard-kpi-label{min-height:2.15em;color:#475569;font-size:.72rem;font-weight:850;line-height:1.18;letter-spacing:.015em;text-transform:uppercase}
.dashboard-kpi-value{margin:.38rem 0 0;color:#0f172a;font-size:clamp(1.65rem,2vw,2.15rem);font-weight:850;line-height:1;letter-spacing:-.045em;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
.dashboard-kpi-caption{margin-top:auto;padding-top:.65rem;color:#718096;font-size:.62rem;font-weight:800;letter-spacing:.04em;text-transform:uppercase}
.dashboard-kpi-progress{height:6px;margin:.62rem 0 0;background:#eef2f7;border-radius:999px;overflow:hidden}
.dashboard-kpi-progress span{display:block;height:100%;border-radius:999px;background:var(--kpi-accent)}
.dashboard-chart-grid{display:grid;grid-template-columns:minmax(320px,.85fr) minmax(460px,1.35fr);gap:.9rem;margin:.1rem 0 .25rem}
.dashboard-chart-card{min-width:0;background:#fff;border:1px solid #e2e8f0;border-radius:14px;box-shadow:0 5px 18px rgba(15,23,42,.045);padding:1rem 1.05rem .8rem;overflow:hidden}
.dashboard-chart-head{display:flex;align-items:flex-start;justify-content:space-between;gap:1rem;margin:0 0 .35rem}
.dashboard-chart-kicker{font-size:.59rem;font-weight:900;letter-spacing:.085em;color:#64748b;text-transform:uppercase}
.dashboard-chart-title{margin:.2rem 0 0;font-size:.86rem;font-weight:900;color:#111827;text-transform:uppercase}
.dashboard-chart-note{margin:.25rem 0 0;color:#64748b;font-size:.68rem;line-height:1.35}
.dashboard-empty{min-height:205px;display:flex;align-items:center;justify-content:center;text-align:center;color:#64748b;font-size:.78rem;font-weight:700;border:1px dashed #d8dee8;border-radius:10px;background:#f8fafc}
@media(max-width:1100px){.dashboard-kpi-grid{grid-template-columns:repeat(2,minmax(0,1fr))}.dashboard-chart-grid{grid-template-columns:1fr}}
@media(max-width:650px){.dashboard-kpi-grid{grid-template-columns:1fr;gap:.65rem}.dashboard-kpi{min-height:126px;padding:.9rem}.dashboard-kpi-label{min-height:0}.dashboard-kpi-value{font-size:1.75rem}}
@media(max-width:900px){.section-title{font-size:1.14rem!important}.api-grid{grid-template-columns:repeat(2,minmax(0,1fr))}}

</style>
""",
    unsafe_allow_html=True,
)

def ncode(s): return s.astype('string').fillna('').str.strip().str.replace(r'\.0$','',regex=True).str.zfill(8)
def naddr(s): return s.astype('string').fillna('').str.strip().str.replace(r'\s+',' ',regex=True).str.upper()
def pnum(v):
 if v is None or pd.isna(v):return 0.0
 s=str(v).strip().replace('R$','').replace(' ','')
 if not s:return 0.0
 neg=s.startswith('-');s=s.lstrip('+-')
 if ',' in s and '.' in s:s=s.replace('.','').replace(',','.') if s.rfind(',')>s.rfind('.') else s.replace(',','')
 elif ',' in s:s=s.replace(',','.')
 try:x=float(s);return -x if neg else x
 except:return 0.0
def nums(s):return s.apply(pnum).astype(float)
def fn(v,d=3):return f'{float(v):,.{d}f}'.replace(',','X').replace('.',',').replace('X','.')
def brl(v):return f'R$ {float(v):,.2f}'.replace(',','X').replace('.',',').replace('X','.')

def eligible_addresses_from_frame(endereco_df):
 if endereco_df is None or endereco_df.empty:return []
 addresses=sorted([x for x in naddr(endereco_df.iloc[:,3]).unique() if x])
 blocked={naddr(pd.Series([x])).iloc[0] for x in ESTOQUE_ENDERECOS_NAO_DISPONIVEIS}
 return [x for x in addresses if x not in blocked]

def build_db(an,end,eligible):
 a=an.iloc[:,[0,3,7,10]].copy();a.columns=['codigo','descricao','qtd_analitico','valor_k'];a['codigo']=ncode(a.codigo);a['descricao']=a.descricao.astype('string').fillna('').str.strip();a['qtd_analitico']=nums(a.qtd_analitico);a['valor_k']=nums(a.valor_k)
 a=a.groupby('codigo',as_index=False).agg(descricao=('descricao','first'),qtd_analitico=('qtd_analitico','sum'),valor_k=('valor_k','sum'));_den=a['qtd_analitico'].where(a['qtd_analitico'].abs()>1e-12);a['valor_unitario']=(a['valor_k']/_den).fillna(0.0)
 e=end.iloc[:,[0,3,7]].copy();e.columns=['codigo','endereco','quantidade'];e['codigo']=ncode(e.codigo);e['endereco']=naddr(e.endereco);e['quantidade']=nums(e.quantidade);e=e.groupby(['codigo','endereco'],as_index=False).quantidade.sum(); ap=set(naddr(pd.Series(eligible)).tolist());e['apto']=e.endereco.isin(ap)
 saldo=e[e.apto].groupby('codigo',as_index=False).quantidade.sum().rename(columns={'quantidade':'saldo_apto'});d=a.merge(saldo,on='codigo',how='outer');d.saldo_apto=d.saldo_apto.fillna(0.0);d.qtd_analitico=d.qtd_analitico.fillna(0.0);d.valor_k=d.valor_k.fillna(0.0);d.valor_unitario=d.valor_unitario.fillna(0.0);d.descricao=d.descricao.fillna('SEM DESCRIÇÃO NO ESTOQUE ANALÍTICO');d['valor_total']=d.saldo_apto*d.valor_unitario;act=(d.saldo_apto>0)&(d.valor_unitario>0);d['classificacao_r_un']=pd.NA;d['classificacao_r_total']=pd.NA;d.loc[act,'classificacao_r_un']=d.loc[act].valor_unitario.rank(method='first',ascending=False).astype(int);d.loc[act,'classificacao_r_total']=d.loc[act].valor_k.rank(method='first',ascending=False).astype(int);d=d.sort_values(['classificacao_r_total','codigo'],na_position='last').reset_index(drop=True);e=e.merge(d[['codigo','valor_unitario']],on='codigo',how='left');return d,e

def nextdoc():
 return central_data.next_inventory_document()
def _counted_cycles(code):
 return int(st.session_state.cycles.get(str(code),0) or 0)

def _open_inventory_codes():
 return {
  str(r.get('codigo'))
  for inv in st.session_state.inventories.values()
  if inv.get('status')!='FECHADO'
  for r in inv.get('rows',[])
 }

def _eligible_priority_frame(db,include_open=False):
 eligible=db[(db.saldo_apto>0)&(db.valor_unitario>0)].copy()
 eligible['cc']=eligible.codigo.astype(str).map(_counted_cycles)
 if not include_open:
  busy=_open_inventory_codes()
  eligible=eligible[~eligible.codigo.astype(str).isin(busy)]
 return eligible

def cycle():
 if st.session_state.db is None:return 1
 frame=_eligible_priority_frame(st.session_state.db,include_open=True)
 return (int(frame.cc.min()) if not frame.empty else 0)+1

def select_products(db,n,urgent_codes=None):
 """Menor número de ciclos primeiro; urgências podem antecipar a fila."""
 available=_eligible_priority_frame(db)
 # A constatação de inconsistência tem prioridade mesmo com saldo de sistema zero.
 codes=set(db.codigo.astype(str))
 busy=_open_inventory_codes()
 urgent=[str(c) for c in dict.fromkeys(urgent_codes or []) if str(c) in codes and str(c) not in busy]
 ordinary=available[~available.codigo.astype(str).isin(urgent)].copy()
 selected=[]
 # Alterna o ranking unitário e o ranking de valor total em cada faixa de ciclos.
 for count in sorted(ordinary.cc.unique()):
  group=ordinary[ordinary.cc==count]
  by_unit=group.sort_values(['classificacao_r_un','codigo'],na_position='last')
  by_total=group.sort_values(['classificacao_r_total','codigo'],na_position='last')
  unit_codes=by_unit.codigo.astype(str).tolist()
  total_codes=by_total.codigo.astype(str).tolist()
  while (unit_codes or total_codes) and len(selected)<int(n):
   for ranking in (unit_codes,total_codes):
    while ranking:
     code=ranking.pop(0)
     if code not in selected:
      selected.append(code)
      break
    if len(selected)>=int(n):break
  if len(selected)>=int(n):break
 # As constatações de inconsistência não consomem as vagas do ciclo regular.
 return db[db.codigo.astype(str).isin(selected+urgent)].copy()

def make_rows(sel,pos):
 rows=[]
 for _,p in sel.iterrows():
  for _,r in pos[(pos.codigo==p.codigo)&pos.apto].iterrows():
   rows.append({
    'id':uuid.uuid4().hex[:12],
    'codigo':str(p.codigo),
    'descricao':str(p.descricao),
    'endereco':str(r.endereco),
    'qtd_sistema':float(r.quantidade),
    'valor_unitario':float(p.valor_unitario),
    'contagens':[],
    'status':'PENDENTE',
    'contagem_final':None,
    'resultado_final':'',
    'comentario_final':'SC',
    'motivo_divergencia':'',
    'tratativa_divergencia':'',
    'observacao_tratativa':'',
    'ajuste_necessario_qtd':0.0,
    'ajuste_autorizado':False,
    'ajuste_qtd':0.0,
    'acao_protheus':'',
    'ajuste_autorizado_por':'',
    'ajuste_autorizado_em':'',
   })
 return rows

def register_divergence_treatment(r,q,motivo,tratamento,observacao=''):
 stamp=now_local()
 ajuste=float(q)-float(r['qtd_sistema'])
 autorizado=str(tratamento).upper()=='AJUSTAR ESTOQUE'
 r['contagem_final']=float(q)
 r['motivo_divergencia']=str(motivo or '').strip()
 r['tratativa_divergencia']=str(tratamento or '').strip().upper()
 r['observacao_tratativa']=str(observacao or '').strip()
 r['ajuste_necessario_qtd']=ajuste
 r['ajuste_autorizado']=autorizado
 r['ajuste_qtd']=ajuste if autorizado else 0.0
 r['acao_protheus']='ENTRADA' if autorizado and ajuste>0 else 'SAÍDA' if autorizado and ajuste<0 else ''
 r['ajuste_autorizado_por']=_session_operator() if autorizado else ''
 r['ajuste_autorizado_em']=stamp.isoformat() if autorizado else ''
 r['resultado_final']='AJUSTE AUTORIZADO' if autorizado else 'ENCERRADO SEM AJUSTE'
 r['comentario_final']=r['observacao_tratativa'] or r['motivo_divergencia'] or 'SC'
 r['status']='FINALIZADO'

def _divergent_pending_rows(inv):
 return [
  r for r in inv.get('rows',[])
  if r.get('status')!='FINALIZADO'
  and r.get('contagens')
  and abs(diff(r,last(r)))>1e-9
 ]

def _render_divergence_treatment(r,q,key_prefix):
 motivo=st.selectbox(
  'Motivo da divergência',
  ['SELECIONE...']+DIVERGENCE_REASONS,
  key=f'{key_prefix}_motivo_{r["id"]}',
 )
 tratamento=st.radio(
  'Tratativa',
  DIVERGENCE_TREATMENTS,
  horizontal=True,
  key=f'{key_prefix}_trat_{r["id"]}',
 )
 observacao=st.text_input(
  'Observação da tratativa (opcional)',
  key=f'{key_prefix}_obs_{r["id"]}',
 )
 if tratamento=='AJUSTAR ESTOQUE':
  _aj=diff(r,q)
  _acao='ENTRADA' if _aj>0 else 'SAÍDA' if _aj<0 else 'SEM AJUSTE'
  st.caption(
   f'Ajuste proposto: {_acao} · {abs(_aj):,.3f} unidade(s) · '
   f'Saldo sistema {fn(r["qtd_sistema"])} → físico {fn(q)}'
  )
 if st.button('REGISTRAR TRATATIVA',key=f'{key_prefix}_final_{r["id"]}',type='primary',use_container_width=True):
  if motivo=='SELECIONE...':
   st.error('Selecione o motivo da divergência antes de concluir.')
   return False
  register_divergence_treatment(r,q,motivo,tratamento,observacao)
  return True
 return False

def persist_inv(inv):
 doc=str(inv.get('documento') or '').strip()
 if not doc:return False
 st.session_state.inventories[doc]=inv
 save('inventories',st.session_state.inventories)
 try:
  central_data.save_inventory_document(doc,inv)
  st.session_state.pop('_operational_persistence_error',None)
  return True
 except Exception as exc:
  st.session_state['_operational_persistence_error']=str(exc)
  return False

def delete_open_inventory(documento):
 doc=str(documento or '').strip()
 inv=st.session_state.inventories.get(doc)
 if _session_profile()!='Gestor' or not inv or inv.get('status')=='FECHADO':
  st.session_state['_inventory_delete_error']='EXCLUSÃO PERMITIDA SOMENTE AO GESTOR PARA INVENTÁRIOS EM ABERTO.'
  return False
 try:
  central_data.delete_inventory_document(doc,_session_auth_token())
  st.session_state.inventories.pop(doc,None)
  save('inventories',st.session_state.inventories)
  if st.session_state.get('selected')==doc:st.session_state.selected=None
  st.session_state.pop('delete_inv_target',None)
  st.session_state.pop('_inventory_delete_error',None)
  st.session_state['_inventory_delete_success']=doc
  return True
 except Exception as exc:
  message=str(exc)
  if 'SESSION_INVALID_OR_EXPIRED' in message or 'AUTH_REQUIRED' in message:
   st.session_state.pop('_setta_auth_user',None)
  st.session_state['_inventory_delete_error']=message
  return False


def forward_problem_to_delivery(report_id):
 try:
  result=central_data.forward_inventory_problem_to_delivery(report_id,_session_auth_token())
  updated=result.get('report')
  if isinstance(updated,dict):
   st.session_state.reports[str(report_id)]=updated
   save('reports',st.session_state.reports)
  st.session_state.pop('_inventory_delivery_sync_error',None)
  return True
 except Exception as exc:
  st.session_state['_inventory_delivery_sync_error']=str(exc)
  return False


def _can_confirm_adjustment():
 # É necessária uma sessão do OperaHub válida, mesmo com leitura pública habilitada.
 return bool(_auth_user() and _session_auth_token()
             and _session_profile() in {'Gestor','Almoxarifado','Operador'})


def confirm_adjustment(documento,item_id,realizado,observacao):
 if not _can_confirm_adjustment():
  st.error('AUTENTIQUE UM OPERADOR DO ALMOXARIFADO PARA CONFIRMAR A EXECUÇÃO.')
  return False
 try:
  result=central_data.confirm_inventory_adjustment(
   documento,item_id,realizado,observacao,_session_auth_token(),
  )
  updated=result.get('document')
  if not isinstance(updated,dict):raise ValueError('RESPOSTA_SEM_DOCUMENTO')
  st.session_state.inventories[str(documento)]=updated
  save('inventories',st.session_state.inventories)
  return True
 except Exception as exc:
  st.error(f'NÃO FOI POSSÍVEL REGISTRAR A VALIDAÇÃO: {exc}')
  return False


def addcount(r,q,cm,stage):
 stamp=now_local()
 r['contagens'].append({
  'etapa':stage,
  'quantidade':float(q),
  'comentario':cm.strip() if cm.strip() else 'SC',
  'data':stamp.strftime('%d/%m/%Y %H:%M:%S'),
  'timestamp':stamp.isoformat(),
  'usuario':_session_operator(),
  'perfil':_session_profile(),
 })
 if stage=='1ª CONTAGEM':
  r['status']='CONTADO'

def last(r):return r['contagens'][-1]['quantidade'] if r['contagens'] else None
def diff(r,q):return float(q)-float(r['qtd_sistema'])
def divergencia_valor(r,q):return diff(r,q)*float(r['valor_unitario'])
def divergência(r,q):return abs(divergencia_valor(r,q))
def sev(v):return 'BAIXO' if v<=100 else 'MÉDIO' if v<=1000 else 'ALTO'

def close_inv(inv):
 doc=str(inv.get('documento') or '').strip()
 pendentes=[r for r in inv.get('rows',[]) if r.get('status')!='FINALIZADO']
 if pendentes:
  st.session_state['_inventory_close_error']=(
   f'{len(pendentes)} item(ns) ainda precisam ser concluídos antes do encerramento.'
  )
  return False
 st.session_state.pop('_inventory_close_error',None)
 before_inv=copy.deepcopy(inv)
 before_cycles=copy.deepcopy(st.session_state.cycles)
 before_reports=copy.deepcopy(st.session_state.reports)

 try:
  for r in inv['rows']:
   if r['contagem_final'] is None:r['contagem_final']=last(r)
   if not r['resultado_final']:r['resultado_final']='ENCERRADO PELO GESTOR'
   r['status']='FINALIZADO'

  cycle_codes=[]
  if not inv.get('ciclo_marcado'):
   cycle_codes=sorted({str(r['codigo']) for r in inv['rows'] if r['contagens']})
   for code in cycle_codes:
    st.session_state.cycles[code]=int(st.session_state.cycles.get(code,0))+1
   inv['ciclo_marcado']=True

  inv['status']='FECHADO'
  inv['encerrado_em']=now_local().isoformat()
  inv['encerrado_por']=_session_operator()

  report_updates={}
  for rid,rep in st.session_state.reports.items():
   if rep.get('status')=='ABERTO' and any(
    str(r['codigo'])==str(rep.get('codigo'))
    and str(r['endereco'])==str(rep.get('endereco'))
    for r in inv['rows']
   ):
    rep['status']='ENCERRADO'
    rep['inventario_doc']=doc
    rep['encerrado_em']=now_local().isoformat()
    rep['encerrado_por']=_session_operator()
    report_updates[str(rid)]=rep

  central_data.close_inventory_atomic(
   doc,
   inv,
   cycle_codes,
   report_updates,
   _session_auth_token(),
  )

  st.session_state.inventories[doc]=inv
  save('inventories',st.session_state.inventories)
  save('cycles',st.session_state.cycles)
  save('reports',st.session_state.reports)
  st.session_state.pop('_operational_persistence_error',None)
  return True
 except Exception as exc:
  st.session_state.inventories[doc]=before_inv
  st.session_state.cycles=before_cycles
  st.session_state.reports=before_reports
  message=str(exc)
  if 'SESSION_INVALID_OR_EXPIRED' in message or 'AUTH_REQUIRED' in message:
   st.session_state.pop('_setta_auth_user',None)
  st.session_state['_operational_persistence_error']=message
  return False


@st.cache_data(show_spinner=False,ttl=3600,max_entries=4)
def _cached_inventory_snapshot(analitico_token,endereco_token):
 # Os tokens entram na chave do cache. Quando qualquer fonte muda,
 # somente a nova combinação é processada novamente.
 an,_=central_data.download_source_frame('analitico',analitico_token,header=1)
 en,_=central_data.download_source_frame('endereco',endereco_token,header=1)
 eligible=eligible_addresses_from_frame(en)
 db,pos=build_db(an,en,eligible)
 return an,en,db,pos,eligible

def _load_inventory_fallback():
 # Primeiro tenta o cache local, que é imediato. Firestore só é consultado
 # se não houver snapshot local utilizável.
 try:
  local_db=load('db')
  local_pos=load('pos')
  local_eligible=load('eligible',[]) or []
  if isinstance(local_db,pd.DataFrame) and isinstance(local_pos,pd.DataFrame):
   st.session_state.db=local_db
   st.session_state.pos=local_pos
   st.session_state.eligible=list(local_eligible)
   return True
 except Exception:
  pass

 try:
  fsdb=firestore_load_db()
  fspos=firestore_load_pos()
  fselig=firestore_load_eligible()
  if isinstance(fsdb,pd.DataFrame) and isinstance(fspos,pd.DataFrame):
   st.session_state.db=fsdb
   st.session_state.pos=fspos
   st.session_state.eligible=list(fselig or [])
   return True
 except Exception:
  pass
 return False

def _save_local_inventory_snapshot():
 # Contingência local sem custo de rede. A persistência remota da base
 # consolidada não é necessária porque ela é reconstruível pela Central.
 try:
  save('db',st.session_state.db)
  save('pos',st.session_state.pos)
  save('eligible',st.session_state.eligible)
 except Exception:
  pass

def sync_central_inventory(force=False):
 try:
  if force:
   central_data.bundle_state.clear()
   central_data.sync_state.clear()
   central_data.download_source_frame.clear()
   _cached_inventory_snapshot.clear()

  bundle=central_data.bundle_state()
  states=central_data.sync_state()
  st.session_state['_central_inventory_bundle']=bundle
  st.session_state['_central_inventory_states']=states

  metas={}
  tokens={}
  for key in ('analitico','endereco'):
   meta=bundle.get(key) or {}
   if not bool(meta.get('available')):
    raise RuntimeError(f'FONTE {key.upper()} NÃO DISPONÍVEL NA CENTRAL.')
   metas[key]=meta
   tokens[key]=central_data.source_token(meta)

  session_token=tokens['analitico']+'||'+tokens['endereco']

  # Se esta sessão já está com a mesma versão, não faz rede/processamento.
  if (
   not force
   and st.session_state.get('_central_inventory_session_token')==session_token
   and isinstance(st.session_state.get('db'),pd.DataFrame)
   and isinstance(st.session_state.get('pos'),pd.DataFrame)
  ):
   return False

  # Depois de uma falha, não repete o processamento pesado a cada clique.
  # Uma nova versão ou o botão REPROCESSAR libera nova tentativa.
  if (
   not force
   and st.session_state.get('_central_inventory_failed_token')==session_token
  ):
   if st.session_state.get('db') is None or st.session_state.get('pos') is None:
    _load_inventory_fallback()
   return False

  needs_commit=False
  for key in ('analitico','endereco'):
   state=states.get(key) or {}
   if (
    str(state.get('version_token') or '')!=tokens[key]
    or str(state.get('status') or '').upper()!='ATUALIZADO'
   ):
    needs_commit=True

  an,en,db,pos,eligible=_cached_inventory_snapshot(
   tokens['analitico'],tokens['endereco']
  )
  st.session_state.an_df=an
  st.session_state.en_df=en
  st.session_state.db=db
  st.session_state.pos=pos
  st.session_state.eligible=list(eligible)
  st.session_state['_central_inventory_session_token']=session_token
  st.session_state.pop('_central_inventory_failed_token',None)
  st.session_state.pop('_central_inventory_error',None)

  # Mantém apenas contingência local. Não grava milhares de linhas no
  # Firestore durante o carregamento normal da interface.
  _save_local_inventory_snapshot()

  if needs_commit:
   for key,frame in (('analitico',an),('endereco',en)):
    meta=metas[key]
    central_data.commit_sync(
     key,
     tokens[key],
     meta.get('last_update_at'),
     len(frame),
     status='ATUALIZADO'
    )

  st.session_state['_central_inventory_success']='ANALÍTICO · ENDEREÇO'
  return True
 except Exception as exc:
  st.session_state['_central_inventory_error']=str(exc)
  try:
   bundle=central_data.bundle_state()
   token_a=central_data.source_token(bundle.get('analitico') or {})
   token_e=central_data.source_token(bundle.get('endereco') or {})
   st.session_state['_central_inventory_failed_token']=token_a+'||'+token_e
  except Exception:
   pass
  if st.session_state.get('db') is None or st.session_state.get('pos') is None:
   _load_inventory_fallback()
  return False

def render_api_monitor():
 section_band('01 · FONTES','ACOMPANHAMENTO DE API')
 try:
  bundle=central_data.bundle_state();states=central_data.sync_state()
 except Exception as exc:
  st.warning(f'CENTRAL INDISPONÍVEL: {exc}');return
 cols=st.columns(2)
 for col,key,label in zip(cols,('analitico','endereco'),('ANALÍTICO','ENDEREÇO')):
  meta=bundle.get(key) or {};state=states.get(key) or {}
  status=str(state.get('status') or ('AGUARDANDO' if meta else 'INDISPONÍVEL')).upper()
  version=f"V{int(meta.get('version') or 0)}"
  when=central_data.format_dt(state.get('synced_at') or meta.get('last_update_at'))
  rows=int(state.get('rows_count') or meta.get('rows_count') or 0)
  col.markdown(f'<div style="background:#fff;border:1px solid #e5e8ee;border-radius:14px;padding:1rem;box-shadow:0 4px 16px rgba(15,23,42,.045)"><div style="font-size:.7rem;font-weight:900;color:#64748b">{label}</div><div style="font-size:1rem;font-weight:900;margin:.25rem 0">{status}</div><div style="font-size:.72rem;color:#667085">{version} · {rows:,} REGISTROS · {when}</div></div>',unsafe_allow_html=True)
 if st.session_state.get('_central_inventory_success'):st.success('FONTES ATUALIZADAS · '+str(st.session_state.pop('_central_inventory_success')))
 if st.session_state.get('_central_inventory_error'):st.warning(str(st.session_state.get('_central_inventory_error')))
 if st.button('REPROCESSAR FONTES',use_container_width=True,key='reprocess_inventory_sources'):
  if sync_central_inventory(force=True):st.rerun()



# Sidebar — espelho estrutural do Controle de NFs
_INV_NAV_PAGES=['Dashboard','Inventário Rotativo','Banco de Dados','Registro','Reportar Inconsistências','Retornos de Produção','Configurações']

def _visible_inventory_pages():
 return ['Retornos de Produção'] if _session_profile()=='Produção' else _INV_NAV_PAGES


def _set_inventory_page(target):
 if target in _visible_inventory_pages():
  st.session_state.section=target
  _setta_close_sidebar()

def _current_inventory_page():
 value=str(st.session_state.get('section') or 'Dashboard')
 if value not in _visible_inventory_pages():
  value='Retornos de Produção' if _session_profile()=='Produção' else 'Dashboard'
  st.session_state.section=value
 return value

with st.sidebar:
 st.markdown(
  '<div class="sidebar-brand">'
  '<div class="sidebar-brand-title">GESTÃO DE ESTOQUE</div>'
  '<div class="sidebar-brand-sub">CONTROLE OPERACIONAL SETTA</div>'
  '</div>'
  '<div class="sidebar-section-label">NAVEGAÇÃO</div>',
  unsafe_allow_html=True,
 )

 _current=_current_inventory_page()
 for _nav_index,_nav_page in enumerate(_visible_inventory_pages()):
  st.button(
   str(_nav_page).upper(),
   key=f'setta_nav_{_nav_index}',
   type='primary' if _nav_page==_current else 'secondary',
   use_container_width=True,
   on_click=_set_inventory_page,
   args=(_nav_page,),
  )

 _sidebar_user=_auth_user()
 if _sidebar_user:
  _sidebar_name=str(_sidebar_user.get('full_name') or _sidebar_user.get('username') or 'USUÁRIO SETTA')
  st.markdown(
   '<div class="sidebar-divider"></div>'
   '<div class="sidebar-section-label">SESSÃO</div>'
   '<div class="sidebar-status-card">'
   f'<div class="sidebar-status-name">{_sidebar_name}</div>'
   f'<div class="sidebar-status-meta">{_session_profile().upper()}</div>'
   '</div>',
   unsafe_allow_html=True,
  )
  if st.button('SAIR',key='setta_auth_logout',use_container_width=True):
   try:
    setta_auth.logout(central_data.supabase_key(),_session_auth_token())
   except Exception:
    pass
   st.session_state.pop('_setta_auth_user',None)
   st.rerun()

 _sidebar_status_slot=st.empty()
 _sidebar_status_slot.markdown(
  '<div class="sidebar-divider"></div><div class="sidebar-section-label">STATUS GERAL</div>'
  '<div class="sidebar-status-card"><div class="sidebar-status-name">GESTÃO DE ESTOQUE</div>'
  '<div class="sidebar-status-value status-warning">ATUALIZANDO</div>'
  '<div class="sidebar-status-meta">SINCRONIZANDO FONTES...</div></div>',
  unsafe_allow_html=True,
 )



if st.session_state.get('_operational_persistence_error'):
 st.error(
  'PERSISTÊNCIA OPERACIONAL INDISPONÍVEL · '
  + str(st.session_state.get('_operational_persistence_error'))
 )

with st.container(key='setta_top_controls'):
 st.button(
  '☰',
  key='setta_drawer_toggle',
  help='Abrir/fechar menu',
  use_container_width=True,
  on_click=_setta_toggle_sidebar,
 )

_main_logo=logo_uri()
_logo_html=(f'<img src="{_main_logo}" alt="SETTA">' if _main_logo else '<div style="font-size:2rem;font-weight:800;color:#202124">SETTA</div>')
st.markdown(f'<div class="setta-logo-card">{_logo_html}</div>',unsafe_allow_html=True)
st.markdown('<h1 class="app-title">GESTÃO DE ESTOQUE | SETTA</h1>',unsafe_allow_html=True)
st.markdown('<p class="app-sub">INVENTÁRIO ROTATIVO • ACURÁCIA • CONTAGENS • HISTÓRICO</p>',unsafe_allow_html=True)

_render_setta_auth_gate()

sync_central_inventory(force=False)

_sidebar_bundle=st.session_state.get('_central_inventory_bundle') or {}
_sidebar_states=st.session_state.get('_central_inventory_states') or {}
_sidebar_docs_total=2
_sidebar_docs_ok=0
_sidebar_has_error=False
_sidebar_latest=None
for _source_key in ('analitico','endereco'):
 _meta=_sidebar_bundle.get(_source_key) or {}
 _state=_sidebar_states.get(_source_key) or {}
 _state_status=str(_state.get('status') or '').upper()
 if bool(_meta.get('available')) and _state_status=='ATUALIZADO':_sidebar_docs_ok+=1
 if _state_status=='ERRO':_sidebar_has_error=True
 _raw_when=str(_state.get('synced_at') or _meta.get('last_update_at') or '').strip()
 if _raw_when:
  _stamp=pd.to_datetime(_raw_when,errors='coerce',utc=True)
  if not pd.isna(_stamp) and (_sidebar_latest is None or _stamp>_sidebar_latest):_sidebar_latest=_stamp

_sidebar_last_update=central_data.format_dt(_sidebar_latest.isoformat()) if _sidebar_latest is not None else '—'
if _sidebar_has_error:
 _sidebar_value='ERRO';_sidebar_status_class='status-error'
elif _sidebar_docs_ok==_sidebar_docs_total:
 _sidebar_value='ATUALIZADO';_sidebar_status_class='status-ok'
else:
 _sidebar_value='ATENÇÃO';_sidebar_status_class='status-warning'

_sidebar_status_slot.markdown(
 f'<div class="sidebar-divider"></div><div class="sidebar-section-label">STATUS GERAL</div>'
 f'<div class="sidebar-status-card"><div class="sidebar-status-name">GESTÃO DE ESTOQUE</div>'
 f'<div class="sidebar-status-value {_sidebar_status_class}">{_sidebar_value}</div>'
 f'<div class="sidebar-status-meta"><div>ÚLTIMA ATUALIZAÇÃO: {_sidebar_last_update}</div>'
 f'<div>QNT DE DOCUMENTOS: {_sidebar_docs_ok}/{_sidebar_docs_total}</div></div></div>',
 unsafe_allow_html=True,
)

active=_current_inventory_page()

# Dashboard
if active=='Dashboard':
 st.markdown('<div class="section-title">DASHBOARD OPERACIONAL</div>',unsafe_allow_html=True)
 section_band('01 · VISÃO GERAL','INDICADORES DO ESTOQUE')
 if st.session_state.db is None:st.info('Aguardando sincronização automática das fontes ANALÍTICO e ENDEREÇO pela API.')
 else:
  db=st.session_state.db;items=int((db.saldo_apto>0).sum());valor_apto=float(db.valor_total.sum());rr=[r for x in st.session_state.inventories.values() for r in x['rows']];cnt=[r for r in rr if r['contagens']];div=[r for r in cnt if abs(diff(r,last(r)))>1e-9]
  qtd_cnt=len(cnt);qtd_div=len(div);acc_itens=(100-(qtd_div/items*100)) if items else 100.0;acc_pos=(100-(qtd_div/qtd_cnt*100)) if qtd_cnt else 100.0
  dashboard_kpi_grid(items,valor_apto,qtd_cnt,qtd_div,acc_itens,acc_pos)
  topic_divider();section_band('02 · INDICADORES','ACOMPANHAMENTO DOS INVENTÁRIOS')
  import altair as alt
  c_left,c_right=st.columns([1,1.45],gap='large')
  with c_left:
   with st.container(border=True):
    st.markdown('**SITUAÇÃO DAS POSIÇÕES CONTADAS**')
    if qtd_cnt:
     status_df=pd.DataFrame({
      'Status':['Sem divergência','Com divergência'],
      'Quantidade':[max(qtd_cnt-qtd_div,0),qtd_div],
     })
     donut=alt.Chart(status_df).mark_arc(innerRadius=66,outerRadius=96).encode(
      theta=alt.Theta('Quantidade:Q',stack=True),
      color=alt.Color('Status:N',scale=alt.Scale(
       domain=['Sem divergência','Com divergência'],range=['#16a34a','#ef4444']
      ),legend=None),
      tooltip=[alt.Tooltip('Status:N',title='Status'),alt.Tooltip('Quantidade:Q',title='Posições')],
     )
     center=alt.Chart(pd.DataFrame({'Total':[qtd_cnt]})).mark_text(
      fontSize=27,fontWeight='bold',color='#111827'
     ).encode(text=alt.Text('Total:Q',format=',d'))
     st.altair_chart((donut+center).properties(height=235,background='transparent').configure_view(stroke=None),use_container_width=True)
     ok_col,div_col=st.columns(2)
     ok_col.metric('SEM DIVERGÊNCIA',max(qtd_cnt-qtd_div,0))
     div_col.metric('COM DIVERGÊNCIA',qtd_div)
    else:
     st.info('NENHUMA CONTAGEM REGISTRADA.')
  with c_right:
   with st.container(border=True):
    st.markdown('**ACURÁCIA DOS ÚLTIMOS INVENTÁRIOS**')
    inv_perf=[]
    _closed_history=[x for x in st.session_state.inventories.values() if x.get('status')=='FECHADO']
    for x in sorted(_closed_history,key=lambda z:z.get('criado_em',''),reverse=True)[:8]:
     total=sum(1 for r in x.get('rows',[]) if r.get('contagens'))
     dv=sum(1 for r in x.get('rows',[]) if r.get('contagens') and abs(diff(r,last(r)))>1e-9)
     if total:
      inv_perf.append({
       'Inventário':x.get('documento',''),
       'Acurácia':max(0.0,100.0-(dv/total*100.0)),
       'Posições':total,'Divergentes':dv,
      })
    if inv_perf:
     perf_df=pd.DataFrame(inv_perf)
     perf_df['Rótulo']=perf_df['Acurácia'].map(lambda v:f'{v:.1f}%')
     bars=alt.Chart(perf_df).mark_bar(cornerRadiusEnd=6,size=24).encode(
      y=alt.Y('Inventário:N',sort=None,axis=alt.Axis(title=None,labelLimit=150,labelFontSize=12)),
      x=alt.X('Acurácia:Q',scale=alt.Scale(domain=[0,112]),
       axis=alt.Axis(title='ACURÁCIA (%)',values=[0,25,50,75,100],grid=True)),
      color=alt.condition('datum.Acurácia >= 98',alt.value('#16a34a'),alt.value('#d97706')),
      tooltip=[alt.Tooltip('Inventário:N',title='Inventário'),
       alt.Tooltip('Acurácia:Q',title='Acurácia (%)',format='.1f'),
       alt.Tooltip('Posições:Q',title='Posições'),
       alt.Tooltip('Divergentes:Q',title='Divergentes')],
     )
     labels=alt.Chart(perf_df).mark_text(align='left',dx=7,fontWeight='bold',fontSize=12).encode(
      y=alt.Y('Inventário:N',sort=None),x='Acurácia:Q',text='Rótulo:N',
     )
     st.altair_chart((bars+labels).properties(height=max(240,48*len(perf_df)),background='transparent').configure_view(stroke=None),use_container_width=True)
    else:
     st.info('NENHUM INVENTÁRIO FECHADO PARA COMPARAÇÃO.')

# Inventory
elif active=='Inventário Rotativo':
 section_band('01 · INVENTÁRIO','CONTROLE E EXECUÇÃO')
 if st.session_state.db is None:st.info('Aguardando sincronização automática da base de estoque pela API.')
 else:
  _,b=st.columns([2,1])
  if b.button(config['new_inventory_text'],type='primary',use_container_width=True):st.session_state.new_inv=True;st.rerun()
  if st.session_state.new_inv:
   with st.container(border=True):
    a,b,c=st.columns(3)
    selection_mode=a.radio('Seleção de materiais',['AUTOMÁTICA','MANUAL'],horizontal=True,key='inventory_selection_mode')
    blind=b.checkbox('Contagem cega',value=cfg['blind_default'])
    c.metric('Ciclo atual',cycle())
    n=10
    manual_codes=[]
    if selection_mode=='AUTOMÁTICA':
     n=st.number_input('Quantidade de produtos distintos',1,500,10,key='inventory_auto_qty')
    else:
     _priority=_eligible_priority_frame(st.session_state.db)
     _urgent_codes={str(rep.get('codigo')) for rep in st.session_state.reports.values() if rep.get('status')=='ABERTO' and rep.get('equipe')=='INVENTÁRIO ROTATIVO'}
     _min_cycle=_priority.cc.min() if not _priority.empty else -1
     _eligible_products=_priority[
      (_priority.cc==_min_cycle)|_priority.codigo.astype(str).isin(_urgent_codes)
     ][['codigo','descricao']].drop_duplicates('codigo').copy()
     _manual_map={
      f'{str(row.codigo)} — {str(row.descricao)}':str(row.codigo)
      for _,row in _eligible_products.iterrows()
     }
     _manual_labels=st.multiselect(
      'Materiais do inventário',
      list(_manual_map.keys()),
      key='inventory_manual_products',
      placeholder='Pesquise pelo código ou descrição',
     )
     manual_codes=[_manual_map[x] for x in _manual_labels]
     st.caption('Inconsistências abertas do Inventário Rotativo continuam sendo incluídas automaticamente.')
    x,y=st.columns(2)
    if x.button('Criar inventário',type='primary',use_container_width=True):
     urgent_codes=sorted({str(rep.get('codigo')) for rep in st.session_state.reports.values() if rep.get('status')=='ABERTO' and rep.get('equipe')=='INVENTÁRIO ROTATIVO'})
     if selection_mode=='MANUAL':
      _priority=_eligible_priority_frame(st.session_state.db)
      _min_cycle=_priority.cc.min() if not _priority.empty else -1
      _priority_codes=set(_priority.loc[_priority.cc==_min_cycle,'codigo'].astype(str))
      _urgent_codes=set(str(c) for c in urgent_codes)
      _selected_codes=[c for c in dict.fromkeys(manual_codes+urgent_codes)
                       if c in _priority_codes or c in _urgent_codes]
      _busy_codes=_open_inventory_codes()
      _selected_codes=[c for c in _selected_codes if c not in _busy_codes]
      sel=st.session_state.db[st.session_state.db.codigo.astype(str).isin(_selected_codes)].copy()
     else:
      sel=select_products(st.session_state.db,n,urgent_codes)
     rows=make_rows(sel,st.session_state.pos)
     if not rows:st.error('Os produtos selecionados não possuem endereços aptos.')
     else:
      try:
       doc=nextdoc()
      except Exception as exc:
       st.error(f'Não foi possível gerar o número do inventário: {exc}')
      else:
       stamp=now_local()
       st.session_state.inventories[doc]={
        'documento':doc,
        'data':stamp.strftime('%d/%m/%Y %H:%M'),
        'responsavel':_session_operator(),
        'perfil_responsavel':_session_profile(),
        'blind_count':blind,
        'ciclo':cycle(),
        'status':'EM CONTAGEM',
        'rows':rows,
        'criado_em':stamp.isoformat(),
        'ciclo_marcado':False,
       }
       if persist_inv(st.session_state.inventories[doc]):
        st.session_state.selected=doc
        st.session_state.new_inv=False
        st.rerun()
       else:
        st.error('Não foi possível salvar o inventário. Tente novamente.')
    if y.button('Cancelar'):st.session_state.new_inv=False;st.rerun()
  def render_inventory_cards(items):
   for inv in items:
    with st.container(border=True):
     s=len(inv['rows']);prod=len({r['codigo'] for r in inv['rows']});a,b,c,d=st.columns([2.2,1.5,1,1]);a.markdown(f'**{inv["documento"]}**');a.caption(f'Ciclo {inv["ciclo"]} · {inv["data"]}');b.write(f'**{inv["status"]}**');c.metric('Produtos',prod);d.metric('Posições',s)
     _show_delete=_session_profile()=='Gestor' and inv.get('status')!='FECHADO'
     _open_col,_delete_col=st.columns(2) if _show_delete else (st.container(),None)
     if _open_col.button('ABRIR',key='op_'+inv['documento'],use_container_width=True):
      st.session_state.selected=inv['documento'];st.rerun()
     if _delete_col is not None and _delete_col.button('EXCLUIR',key='del_'+inv['documento'],use_container_width=True):
      st.session_state.delete_inv_target=inv['documento'];st.rerun()
  all_inv=sorted(st.session_state.inventories.values(),key=lambda x:x.get('criado_em',''),reverse=True)
  open_inv=[x for x in all_inv if x.get('status')!='FECHADO']
  closed_inv=[x for x in all_inv if x.get('status')=='FECHADO']
  tab_open,tab_closed=st.tabs([f'EM ABERTO ({len(open_inv)})',f'FECHADOS ({len(closed_inv)})'])
  with tab_open:
   if open_inv:render_inventory_cards(open_inv)
   else:st.info('Nenhum inventário em aberto.')
  with tab_closed:
   if closed_inv:render_inventory_cards(closed_inv)
   else:st.info('Nenhum inventário fechado.')
  if st.session_state.pop('_inventory_delete_success',None):
   st.success('INVENTÁRIO EXCLUÍDO. AS CONTAGENS FORAM DESCARTADAS.')
  if st.session_state.get('_inventory_delete_error'):
   st.error('NÃO FOI POSSÍVEL EXCLUIR: '+str(st.session_state['_inventory_delete_error']))
  _delete_target=st.session_state.get('delete_inv_target')
  _delete_inv=st.session_state.inventories.get(_delete_target) if _delete_target else None
  if _delete_inv and _delete_inv.get('status')!='FECHADO' and _session_profile()=='Gestor':
   with st.container(border=True):
    st.warning(f'EXCLUIR INVENTÁRIO {_delete_target}? ESTA AÇÃO É IRREVERSÍVEL E DESCARTA TODAS AS CONTAGENS DESTE DOCUMENTO.')
    _typed=st.text_input('DIGITE O NÚMERO DO INVENTÁRIO PARA CONFIRMAR',key='inventory_delete_confirmation',placeholder=_delete_target)
    _yes,_no=st.columns(2)
    if _yes.button('CONFIRMAR EXCLUSÃO',key='confirm_delete_open_inv',type='primary',use_container_width=True,disabled=_typed.strip()!=_delete_target):
     delete_open_inventory(_delete_target)
     st.rerun()
    if _no.button('CANCELAR EXCLUSÃO',key='cancel_delete_open_inv',use_container_width=True):
     st.session_state.pop('delete_inv_target',None)
     st.session_state.pop('_inventory_delete_error',None)
     st.rerun()
  elif _delete_target:
   st.session_state.pop('delete_inv_target',None)
  doc=st.session_state.selected
  if doc in st.session_state.inventories:
   inv=st.session_state.inventories[doc];st.divider();st.markdown(f'### Inventário {doc} — {inv["status"]}')
   prof=_session_profile()
   if prof=='Operador' and inv['status']=='EM CONTAGEM':
    total_pos=len(inv['rows'])
    counted_pos=sum(1 for r in inv['rows'] if r['contagens'])
    pending_pos=total_pos-counted_pos
    p1,p2,p3=st.columns(3)
    p1.metric('POSIÇÕES',total_pos)
    p2.metric('CONTADAS',counted_pos)
    p3.metric('PENDENTES',pending_pos)
    st.progress(counted_pos/total_pos if total_pos else 0.0)

    if total_pos and counted_pos==total_pos:
     st.markdown('#### 1ª CONTAGEM CONCLUÍDA')
     if st.button('FINALIZAR 1ª CONTAGEM',type='primary',use_container_width=True,key='finish_first_count'):
      inv['status']='AGUARDANDO ANÁLISE'
      if persist_inv(inv):
       st.rerun()
      else:
       st.error('Não foi possível finalizar a 1ª contagem.')
     st.success('Todas as posições foram contadas. Finalize acima para encaminhar o inventário à análise do Gestor.')
    else:
     st.info(f'Faltam {pending_pos} posição(ões) para liberar a finalização.')

    for r in inv['rows']:
     if r['contagens']:continue
     with st.container(border=True):
      a,b,c=st.columns([1.1,3,1.3]);a.markdown(f'**{r["codigo"]}**');b.write(f'{r["descricao"]}\n\n**Endereço:** {r["endereco"]}');c.write('**Qtd. sistema:** OCULTA' if inv['blind_count'] else f'**Qtd. sistema:** {fn(r["qtd_sistema"])}');x,y=st.columns([1,2]);q=x.number_input('Contagem',0.0,step=.001,format='%.3f',key='q1_'+r['id']);cm=y.text_input('Comentário (opcional)',key='cm1_'+r['id'])
      if st.button('Salvar contagem',key='sv1_'+r['id'],type='primary'):
       addcount(r,q,cm,'1ª CONTAGEM')
       if persist_inv(inv):st.rerun()
       else:st.error('Não foi possível salvar a contagem.')
   elif prof=='Gestor' and inv['status']=='AGUARDANDO ANÁLISE':
    st.markdown('#### Análise da 1ª contagem')
    # 1ª contagem: igual ao sistema confirma; divergente segue para decisão.
    for r in inv['rows']:
     if not r['contagens'] or r['status']=='FINALIZADO':
      continue
     q=last(r)
     if abs(diff(r,q))<1e-9:
      r['contagem_final']=q;r['resultado_final']='SISTEMA CONFIRMADO';r['status']='FINALIZADO'
    persist_inv(inv)
    divs=[r for r in inv['rows'] if r['contagens'] and r['status']!='FINALIZADO' and abs(diff(r,last(r)))>1e-9]
    if not divs: st.success('Não existem divergências na 1ª contagem. Todos os itens foram confirmados pelo sistema.')
    for r in divs:
     with st.container(border=True):
      q=last(r);a,b,c,d=st.columns(4);a.markdown(f'**{r["codigo"]} / {r["endereco"]}**');a.caption(r.get('descricao',''));b.metric('Sistema',fn(r['qtd_sistema']));c.metric('1ª contagem',fn(q));d.metric('Divergência',signed_brl(divergencia_valor(r,q)));st.caption(f'Divergência de quantidade: {diff(r,q):+,.3f}'.replace(',','X').replace('.',',').replace('X','.').replace('+','+')+f' · Classificação: {sev(abs(divergencia_valor(r,q)))} · Comentário: {r["contagens"][-1]["comentario"]}')
      x,y=st.columns(2)
      if x.button('RECONTAR ESTE ITEM',key='r1_'+r['id'],use_container_width=True):r['status']='RECONTAR';inv['status']='AGUARDANDO RECONTAGEM';persist_inv(inv);st.rerun()
      if y.button('AUDITAR ESTE ITEM',key='a1_'+r['id'],use_container_width=True):r['status']='AUDITORIA';inv['status']='AGUARDANDO AUDITORIA';persist_inv(inv);st.rerun()
      with st.expander('CONCLUIR / TRATAR DIVERGÊNCIA'):
       if _render_divergence_treatment(r,q,'a1'):
        persist_inv(inv);st.rerun()
    x,y,z=st.columns(3)
    if x.button('RECONTAR TODOS OS DIVERGENTES',type='primary',use_container_width=True):
     for r in inv['rows']:
      if r['contagens'] and r['status']!='FINALIZADO' and abs(diff(r,last(r)))>1e-9:r['status']='RECONTAR'
     inv['status']='AGUARDANDO RECONTAGEM';persist_inv(inv);st.rerun()
    if y.button('AUDITAR TODOS OS DIVERGENTES',use_container_width=True):
     for r in inv['rows']:
      if r['contagens'] and r['status']!='FINALIZADO' and abs(diff(r,last(r)))>1e-9:r['status']='AUDITORIA'
     inv['status']='AGUARDANDO AUDITORIA';persist_inv(inv);st.rerun()
    if z.button('ENCERRAR INVENTÁRIO',type='primary',use_container_width=True):
     if close_inv(inv):
       st.session_state.section='Registro' if any(r.get('ajuste_autorizado') for r in inv.get('rows',[])) else 'Inventário Rotativo'
       st.rerun()
     else:st.error(st.session_state.get('_inventory_close_error') or 'Não foi possível encerrar o inventário. A operação foi revertida.')
   elif prof=='Operador' and inv['status']=='AGUARDANDO RECONTAGEM':
    st.markdown('#### Recontagem — itens liberados pelo gestor')
    targets=[r for r in inv['rows'] if r['status']=='RECONTAR']
    for r in targets:
     with st.container(border=True):
      st.markdown(f'**{r["codigo"]} / {r["endereco"]}**');st.write(r['descricao']);st.caption('Histórico: '+' → '.join(f"{x['etapa']}: {fn(x['quantidade'])}" for x in r['contagens']));q=st.number_input('Nova contagem',0.0,step=.001,format='%.3f',key='q2_'+r['id']);cm=st.text_input('Comentário (opcional)',key='cm2_'+r['id'])
      if st.button('Salvar nova contagem',key='sv2_'+r['id'],type='primary'):addcount(r,q,cm,f'{len(r["contagens"])+1}ª CONTAGEM');r['status']='RECONTADA';persist_inv(inv);st.rerun()
    if not any(r['status']=='RECONTAR' for r in inv['rows']):inv['status']='AGUARDANDO DECISÃO';persist_inv(inv);st.rerun()
   elif prof=='Gestor' and inv['status']=='AGUARDANDO DECISÃO':
    st.markdown('#### Avaliação após cada recontagem')
    # Mesma regra para 2ª, 3ª e todas as contagens seguintes.
    candidates=[r for r in inv['rows'] if r['contagens'] and r['status']!='FINALIZADO' and abs(diff(r,last(r)))>1e-9]
    if not candidates: st.success('Não existem itens pendentes de decisão.')
    for r in candidates:
     q2=last(r);q1=r['contagens'][-2]['quantidade'] if len(r['contagens'])>=2 else None
     with st.container(border=True):
      st.markdown(f'**{r["codigo"]} / {r["endereco"]}**');st.caption(r.get('descricao',''));st.write(f'Sistema: **{fn(r["qtd_sistema"])}** · Anterior: **{fn(q1) if q1 is not None else "—"}** · Atual: **{fn(q2)}**')
      if abs(q2-r['qtd_sistema'])<1e-9:
       st.success('Atual igual ao sistema: SISTEMA CONFIRMADO. Nenhuma nova contagem é necessária.');r['contagem_final']=q2;r['resultado_final']='SISTEMA CONFIRMADO';r['status']='FINALIZADO';persist_inv(inv)
      else:
       if q1 is not None and abs(q2-q1)<1e-9:
        st.success('A recontagem confirmou o físico. A divergência com o sistema permanece e agora exige tratativa do Gestor.')
       else:
        st.warning('A divergência permanece. O gestor deve decidir o próximo passo.')
       a,b=st.columns(2)
       if a.button('RECONTAR ESTE ITEM',key='dr_'+r['id'],use_container_width=True):r['status']='RECONTAR';inv['status']='AGUARDANDO RECONTAGEM';persist_inv(inv);st.rerun()
       if b.button('AUDITAR ESTE ITEM',key='da_'+r['id'],use_container_width=True):r['status']='AUDITORIA';inv['status']='AGUARDANDO AUDITORIA';persist_inv(inv);st.rerun()
       with st.expander('CONCLUIR / TRATAR DIVERGÊNCIA'):
        if _render_divergence_treatment(r,q2,'dec'):
         persist_inv(inv);st.rerun()
    x,y,z=st.columns(3)
    if x.button('RECONTAR TODOS OS DIVERGENTES',type='primary',use_container_width=True):
     for r in inv['rows']:
      if r['contagens'] and r['status']!='FINALIZADO' and abs(diff(r,last(r)))>1e-9:r['status']='RECONTAR'
     inv['status']='AGUARDANDO RECONTAGEM';persist_inv(inv);st.rerun()
    if y.button('AUDITAR TODOS OS DIVERGENTES',use_container_width=True):
     for r in inv['rows']:
      if r['contagens'] and r['status']!='FINALIZADO' and abs(diff(r,last(r)))>1e-9:r['status']='AUDITORIA'
     inv['status']='AGUARDANDO AUDITORIA';persist_inv(inv);st.rerun()
    if z.button('ENCERRAR INVENTÁRIO',type='primary',use_container_width=True):
     if close_inv(inv):
       st.session_state.section='Registro' if any(r.get('ajuste_autorizado') for r in inv.get('rows',[])) else 'Inventário Rotativo'
       st.rerun()
     else:st.error('Não foi possível encerrar o inventário. A operação foi revertida.')
   elif prof=='Gestor' and inv['status']=='AGUARDANDO AUDITORIA':
    st.markdown('#### Auditoria / 3ª ou próxima contagem')
    targets=[r for r in inv['rows'] if r['status']=='AUDITORIA']
    for r in targets:
     with st.container(border=True):
      st.markdown(f'**{r["codigo"]} / {r["endereco"]}**');st.write(r['descricao']);st.caption('Histórico: '+' → '.join(f"{x['etapa']}: {fn(x['quantidade'])}" for x in r['contagens']));q=st.number_input('Contagem de auditoria',0.0,step=.001,format='%.3f',key='q3_'+r['id']);cm=st.text_input('Comentário (opcional)',key='cm3_'+r['id'])
      if st.button('Salvar auditoria',key='sv3_'+r['id'],type='primary'):addcount(r,q,cm,'AUDITORIA');r['status']='AUDITADA';persist_inv(inv);st.rerun()
    if not any(r['status']=='AUDITORIA' for r in inv['rows']):inv['status']='AGUARDANDO DECISÃO';persist_inv(inv);st.rerun()
   elif prof=='Gestor' and inv['status']=='FECHADO':
    st.success('Inventário encerrado e salvo no registro.');
    if st.button('Reabrir análise'):inv['status']='AGUARDANDO DECISÃO';persist_inv(inv);st.rerun()

# Database — somente consulta da base tratada publicada pela Central
elif active=='Banco de Dados':
 section_band('01 · BASE CENTRAL','ESTOQUE TRATADO','CONSULTA SOMENTE LEITURA · ATUALIZAÇÃO AUTOMÁTICA VIA API')
 try:
  _derived_status=central_data.derived_status(['estoque_tratado'])
  _treated_meta=_derived_status.get('estoque_tratado') or {}
  if not bool(_treated_meta.get('available')):
   st.info('A base ESTOQUE TRATADO ainda não está disponível na Central de Dados.')
  else:
   _treated_token=f"v{int(_treated_meta.get('version') or 0)}|{_treated_meta.get('processed_at') or _treated_meta.get('last_update_at') or ''}"
   _treated,_treated_download_meta=central_data.download_derived('estoque_tratado',_treated_token)
   _when=central_data.format_dt(_treated_meta.get('processed_at') or _treated_meta.get('last_update_at'))
   _saldo=float(pd.to_numeric(_treated.get('SALDO_DISPONIVEL',pd.Series(dtype=float)),errors='coerce').fillna(0).sum()) if not _treated.empty else 0.0
   _versao=int(_treated_meta.get('version') or 0)

   # Cards próprios desta tela: não dependem de setta_kpi(),
   # evitando qualquer incompatibilidade entre versões durante deploy.
   st.markdown(
    '<div class="dashboard-kpi-grid">'
    '<div class="dashboard-kpi dashboard-kpi--blue">'
    '<div class="dashboard-kpi-top"><span class="dashboard-kpi-eyebrow">BASE TRATADA</span><span class="dashboard-kpi-status-dot"></span></div>'
    '<div class="dashboard-kpi-label">REGISTROS</div>'
    f'<div class="dashboard-kpi-value">{f"{len(_treated):,}".replace(",", ".")}</div>'
    '<div class="dashboard-kpi-caption">ESTOQUE TRATADO</div></div>'
    '<div class="dashboard-kpi dashboard-kpi--green">'
    '<div class="dashboard-kpi-top"><span class="dashboard-kpi-eyebrow">DISPONIBILIDADE</span><span class="dashboard-kpi-status-dot"></span></div>'
    '<div class="dashboard-kpi-label">SALDO DISPONÍVEL</div>'
    f'<div class="dashboard-kpi-value">{fn(_saldo)}</div>'
    '<div class="dashboard-kpi-caption">SOMA DA BASE TRATADA</div></div>'
    '<div class="dashboard-kpi dashboard-kpi--violet">'
    '<div class="dashboard-kpi-top"><span class="dashboard-kpi-eyebrow">CENTRAL DE DADOS</span><span class="dashboard-kpi-status-dot"></span></div>'
    '<div class="dashboard-kpi-label">VERSÃO</div>'
    f'<div class="dashboard-kpi-value">V{_versao}</div>'
    f'<div class="dashboard-kpi-caption">{_when}</div></div>'
    '</div>',
    unsafe_allow_html=True
   )

   topic_divider()
   section_band('02 · RELATÓRIO','ESTOQUE TRATADO')
   _view=_treated.copy()
   _rename={
    'COD_MATERIAL':'CÓDIGO',
    'DESCRICAO':'DESCRIÇÃO',
    'SALDO_EM_ESTOQUE':'SALDO EM ESTOQUE',
    'SALDO_NAO_DISPONIVEL':'SALDO NÃO DISPONÍVEL',
    'SALDO_DISPONIVEL':'SALDO DISPONÍVEL',
   }
   _view=_view.rename(columns=_rename)
   for _col in ['SALDO EM ESTOQUE','SALDO NÃO DISPONÍVEL','SALDO DISPONÍVEL']:
    if _col in _view.columns:_view[_col]=pd.to_numeric(_view[_col],errors='coerce').fillna(0).map(fn)
   _setta_dataframe(_view,use_container_width=True,hide_index=True,height=560)
   st.caption('REGRA DE DISPONIBILIDADE IDÊNTICA AO RELATÓRIO ESTOQUE TRATADO. NÃO HÁ EDIÇÃO, IMPORTAÇÃO OU SELEÇÃO MANUAL NESTA TELA.')
 except Exception as exc:
  st.warning(f'Não foi possível consultar o ESTOQUE TRATADO na Central: {exc}')


# Register
elif active=='Registro':
 section_band('01 · REGISTRO','HISTÓRICO DE INVENTÁRIOS');rows=[]
 for inv in st.session_state.inventories.values():
  if inv['status']!='FECHADO':continue
  for r in inv['rows']:
   _final=r.get('contagem_final')
   _ajuste_necessario=float(r.get('ajuste_necessario_qtd') or (diff(r,_final) if _final is not None else 0))
   rows.append({
    'Documento':inv['documento'],
    'Data':inv['data'],
    'Responsável':inv['responsavel'],
    'Ciclo':inv['ciclo'],
    'ID Interno':r.get('id',''),
    'Código':r['codigo'],
    'Descrição':r['descricao'],
    'Endereço':r['endereco'],
    'Qtd. Sistema':r['qtd_sistema'],
    'Contagens':' | '.join(f"{x['etapa']}: {fn(x['quantidade'])} ({x['comentario']})" for x in r['contagens']),
    'Contagem Final':_final,
    'Resultado':r.get('resultado_final',''),
    'Motivo Divergência':r.get('motivo_divergencia',''),
    'Tratativa':r.get('tratativa_divergencia',''),
    'Ajuste Necessário':_ajuste_necessario,
    'Ajuste Autorizado':'SIM' if r.get('ajuste_autorizado') else 'NÃO',
    'Status Ajuste':r.get('ajuste_status') or ('PENDENTE' if r.get('ajuste_autorizado') else 'NÃO APLICÁVEL'),
    'Ajuste Validado por':r.get('ajuste_validado_por',''),
    'Ajuste Validado em':central_data.format_dt(r.get('ajuste_validado_em')) if r.get('ajuste_validado_em') else '',
    'Observação Ajuste':r.get('ajuste_observacao',''),
    'Qtd. Ajuste':float(r.get('ajuste_qtd') or 0),
    'Ação Protheus':r.get('acao_protheus',''),
    'Autorizado por':r.get('ajuste_autorizado_por',''),
    'Autorizado em':r.get('ajuste_autorizado_em',''),
    'Valor Divergência':divergencia_valor(r,_final) if _final is not None else 0,
   })
 if rows:
  df=pd.DataFrame(rows)
  display_df=df.drop(columns=['ID Interno']).copy()
  display_df['Valor Divergência']=display_df['Valor Divergência'].map(signed_brl)
  _setta_dataframe(display_df,use_container_width=True,hide_index=True)
  export_df=display_df.copy()
  c1,c2,c3=st.columns(3)
  c1.download_button('Exportar Registro em Excel',excel_bytes(export_df,'Registro'),'registro_inventarios.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',use_container_width=True)
  # Só a confirmação SIM dá baixa; NÃO continua listado até execução.
  ajustes=df[df['Ajuste Autorizado'].eq('SIM') & df['Status Ajuste'].ne('REALIZADO')].copy()
  topic_divider()
  section_band('02 · PROTHEUS','AJUSTES PENDENTES DE CONFIRMAÇÃO')
  if not ajustes.empty:
   st.markdown(f"**{len(ajustes)} AJUSTE(S) AGUARDANDO CONFIRMAÇÃO DE EXECUÇÃO**")
   if _can_confirm_adjustment():
    # Confirmação visível antes do relatório, não escondida abaixo da tabela.
    with st.container(border=True):
     st.markdown('#### CONFIRMAR AJUSTE NO PROTHEUS')
     opcoes={
      f"{row['Documento']} | {row['Código']} | {row['Endereço']} | {row['Status Ajuste']}":
      (row['Documento'],row['ID Interno'])
      for _,row in ajustes.iterrows()
     }
     with st.form('confirmacao_ajuste_protheus',clear_on_submit=True):
      alvo=st.selectbox('SELECIONE O MATERIAL / ENDEREÇO',list(opcoes.keys()))
      _doc,_id=opcoes[alvo]
      _selected=ajustes[
       ajustes['Documento'].eq(_doc) & ajustes['ID Interno'].eq(_id)
      ].iloc[0]
      st.markdown(f"**DESCRIÇÃO:** {_selected['Descrição']}")
      _a,_b,_c=st.columns(3)
      _a.metric('SALDO SISTEMA',fn(_selected['Qtd. Sistema']))
      _b.metric('SALDO FÍSICO',fn(_selected['Contagem Final']))
      _c.metric('AJUSTE NECESSÁRIO',fn(_selected['Ajuste Necessário']))
      realizado=st.radio('O AJUSTE FOI LANÇADO NO PROTHEUS?',['SIM','NÃO'],horizontal=True)
      observacao=st.text_input(
       'NÚMERO DO MOVIMENTO / OBSERVAÇÃO',
       placeholder='Se SIM: nº do lançamento. Se NÃO: motivo da pendência.',
      )
      salvar=st.form_submit_button(
       'REGISTRAR CONFIRMAÇÃO',type='primary',use_container_width=True,
      )
     if salvar:
      if not observacao.strip():
       st.error('INFORME O NÚMERO DO LANÇAMENTO OU O MOTIVO DA PENDÊNCIA.')
      elif confirm_adjustment(_doc,_id,realizado=='SIM',observacao):
       st.rerun()
   else:
    st.info('PARA CONFIRMAR AJUSTES, IDENTIFIQUE-SE COM UM USUÁRIO SETTA DO ALMOXARIFADO.')
   st.markdown('#### AJUSTES AINDA PENDENTES')
   pendentes=ajustes[[
    'Documento','Código','Descrição','Endereço',
    'Ajuste Necessário','Ação Protheus','Status Ajuste'
   ]].copy()
   _setta_dataframe(pendentes,use_container_width=True,hide_index=True)
  else:
   st.success('NENHUM AJUSTE PENDENTE DE CONFIRMAÇÃO.')
  if not ajustes.empty:
   carga=ajustes[[
    'Documento','Código','Descrição','Endereço','Qtd. Sistema','Contagem Final',
    'Ação Protheus','Qtd. Ajuste','Ajuste Necessário','Motivo Divergência',
    'Autorizado por','Autorizado em'
   ]].copy()
   carga['QUANTIDADE PROTHEUS']=carga['Qtd. Ajuste'].abs()
   carga['AJUSTE QTD (SINAL)']=carga['Ajuste Necessário']
   carga=carga.rename(columns={
    'Ação Protheus':'AÇÃO PROTHEUS',
    'Qtd. Sistema':'SALDO SISTEMA',
    'Contagem Final':'SALDO FÍSICO',
    'Motivo Divergência':'MOTIVO',
    'Autorizado por':'AUTORIZADO POR',
    'Autorizado em':'AUTORIZADO EM',
   })
   carga=carga[[
    'Documento','Código','Descrição','Endereço','AÇÃO PROTHEUS',
    'QUANTIDADE PROTHEUS','AJUSTE QTD (SINAL)','SALDO SISTEMA','SALDO FÍSICO',
    'MOTIVO','AUTORIZADO POR','AUTORIZADO EM'
   ]]
   c2.download_button('Ajustes Protheus · Excel',excel_bytes(carga,'AJUSTES PROTHEUS'),'ajustes_protheus.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',use_container_width=True)
   _csv=carga.to_csv(index=False,sep=';',decimal=',').encode('utf-8-sig')
   c3.download_button('Ajustes Protheus · CSV',_csv,'ajustes_protheus.csv','text/csv',use_container_width=True)
  else:
   c2.caption('Nenhum ajuste pendente para exportação.')
 else:st.info('Nenhum inventário fechado.')

# Reportar Inconsistências
elif active=='Reportar Inconsistências':
 section_band('01 · INCONSISTÊNCIAS','REGISTRO E ACOMPANHAMENTO')
 if st.session_state.db is None:
  st.warning('Primeiro importe e processe a base na aba Banco de Dados.')
 else:
  if st.button('INICIAR INCONSISTÊNCIA',type='primary',use_container_width=True):st.session_state.new_report=True;st.rerun()
  if 'new_report' not in st.session_state:st.session_state.new_report=False
  if st.session_state.new_report:
   with st.container(border=True):
    st.markdown('#### Novo chamado de inconsistência')
    st.info('O material reportado será incluído automaticamente no próximo Inventário Rotativo.')
    codes=sorted(st.session_state.db['codigo'].astype(str).unique())
    addresses=sorted([str(x) for x in st.session_state.eligible if str(x).strip()])
    a,b=st.columns(2)
    equipe=a.selectbox('Equipe responsável',['INVENTÁRIO ROTATIVO'])
    codigo=b.selectbox('Código do material',options=['']+codes,index=0,help='Digite para pesquisar; a lista sugere correspondências.')
    descricao=''
    if codigo:
     m=st.session_state.db[st.session_state.db.codigo.astype(str)==str(codigo)]
     if not m.empty:descricao=str(m.iloc[0]['descricao'])
    st.caption(f'Descrição: {descricao}' if descricao else 'Selecione o código do material.')
    endereco=st.selectbox('Endereço',options=['']+addresses,index=0,help='Digite para pesquisar entre os endereços habilitados no Banco de Dados.')
    psy=st.text_input('PSY / PROJETO DA INCONSISTÊNCIA',placeholder='Informe os 11 dígitos da PSY para vincular à Gestão de Entregas')
    obs=st.text_area('OBSERVAÇÃO OBRIGATÓRIA',placeholder='INFORME O QUE ACONTECEU...',height=130)
    st.caption('A observação será registrada automaticamente em CAIXA ALTA.')
    x,y=st.columns(2)
    if x.button('SALVAR INCONSISTÊNCIA',type='primary',use_container_width=True):
     obs=obs.strip().upper()
     psy=psy.strip()
     if not codigo or not endereco or not obs or not (len(psy)==11 and psy.isdigit()):
      st.error('Código, endereço, observação e PSY com 11 dígitos são obrigatórios.')
     else:
      stamp=now_local()
      rid=stamp.strftime('%Y%m%d%H%M%S')+'-'+uuid.uuid4().hex[:6].upper()
      report={
       'id':rid,
       'criado_em':stamp.isoformat(),
       'criado_por':_session_operator(),
       'perfil_criador':_session_profile(),
       'equipe':equipe,
       'codigo':str(codigo),
       'descricao':descricao,
       'endereco':str(endereco),
       'psy':psy,
       'observacao':obs,
       'status':'ABERTO',
       'inventario_doc':None,
       'encerrado_em':None,
      }
      if persist_report(report):
       st.session_state.new_report=False
       forward_problem_to_delivery(rid)
       st.rerun()
      else:
       st.error('Não foi possível registrar a inconsistência.')
    if y.button('CANCELAR',use_container_width=True):st.session_state.new_report=False;st.rerun()
  if st.session_state.get('_inventory_delivery_sync_error'):
   st.warning('INCONSISTÊNCIA REGISTRADA, MAS O ENCAMINHAMENTO A ENTREGAS AINDA ESTÁ PENDENTE: '+str(st.session_state['_inventory_delivery_sync_error']))
  reports=sorted(st.session_state.reports.values(),key=lambda x:x.get('criado_em',''),reverse=True)
  abertos=[r for r in reports if r.get('status')=='ABERTO'];encerrados=[r for r in reports if r.get('status')=='ENCERRADO']
  st.divider();st.markdown('### Chamados')
  ta,te=st.tabs([f'NÃO TRATADOS ({len(abertos)})',f'ENCERRADOS ({len(encerrados)})'])
  with ta:
   if not abertos:st.success('Nenhuma inconsistência pendente.')
   for r in abertos:
    with st.container(border=True):
     a,b,c=st.columns([1.2,2.2,1.5]);a.markdown(f'**{r["id"]}**');b.markdown(f'**{r["codigo"]}** · {r["descricao"]}');c.markdown(f'**{r["endereco"]}**')
     st.caption(f'Criado em {r["criado_em"]} · Equipe: {r["equipe"]}')
     st.write(f'**OBSERVAÇÃO:** {r["observacao"]}')
     st.warning('PENDENTE — será incluído no próximo Inventário Rotativo.')
     if r.get('delivery_sync_at'):
      st.success(f"ENCAMINHADO A ENTREGAS · PSY {r.get('psy') or '—'}")
     elif r.get('psy'):
      st.warning('ENCAMINHAMENTO À GESTÃO DE ENTREGAS PENDENTE.')
      if _auth_user() and st.button('TENTAR ENCAMINHAR PARA ENTREGAS',key='delivery_report_'+str(r['id'])):
       if forward_problem_to_delivery(r['id']):st.rerun()
       else:st.rerun()
     else:
      st.caption('REGISTRO ANTIGO SEM PSY: NÃO FOI ENCAMINHADO A ENTREGAS.')
  with te:
   if not encerrados:st.info('Nenhuma inconsistência encerrada.')
   for r in encerrados:
    with st.container(border=True):
     a,b,c=st.columns([1.2,2.2,1.5]);a.markdown(f'**{r["id"]}**');b.markdown(f'**{r["codigo"]}** · {r["descricao"]}');c.markdown(f'**{r["endereco"]}**')
     st.caption(f'Criado em {r["criado_em"]} · Encerrado em {r.get("encerrado_em") or "—"} · Equipe: {r["equipe"]}')
     st.write(f'**OBSERVAÇÃO:** {r["observacao"]}')
     st.success(f'TRATADO NO INVENTÁRIO: {r.get("inventario_doc") or "—"}')

# Produção -> Almoxarifado -> Protheus. Operadores da produção enxergam somente esta tela.
elif active=='Retornos de Produção':
 section_band('01 · RETORNOS','DEVOLUÇÕES DA PRODUÇÃO')
 if not _auth_user():
  st.info('IDENTIFIQUE-SE COM O SEU USUÁRIO SETTA PARA UTILIZAR OS RETORNOS.')
  with st.form('return_login_form'):
   _login=st.text_input('USUÁRIO / E-MAIL')
   _password=st.text_input('SENHA',type='password')
   _submit=st.form_submit_button('ENTRAR',type='primary',use_container_width=True)
  if _submit:
   try:
    if _authenticate_user(_login,_password):st.rerun()
    else:st.error('USUÁRIO OU SENHA INVÁLIDOS.')
   except Exception as exc:st.error(f'ERRO NA AUTENTICAÇÃO: {exc}')
 elif _session_profile() not in {'Produção','Almoxarifado','Gestor'}:
  st.warning('ACESSO AOS RETORNOS RESTRITO A USUÁRIOS AUTORIZADOS DA PRODUÇÃO OU ALMOXARIFADO.')
 else:
  role=_session_profile()
  try:
   returns=central_data.list_production_returns(_session_auth_token())
  except Exception as exc:
   st.error(f'NÃO FOI POSSÍVEL CONSULTAR OS RETORNOS: {exc}')
   returns=[]
  _pending=sum(1 for item in returns if item.get('status')=='PENDENTE')
  _received=sum(1 for item in returns if item.get('status')=='RECEBIDO')
  _closed=sum(1 for item in returns if item.get('status') in ('CONCLUIDO','RECUSADO'))
  _k1,_k2,_k3=st.columns(3)
  _k1.metric('PENDENTES DE RECEBIMENTO',_pending)
  _k2.metric('RECEBIDOS · AGUARDANDO CONCLUSÃO',_received)
  _k3.metric('CONCLUÍDOS / RECUSADOS',_closed)
  if role=='Produção':
   section_band('02 · SOLICITAR','INFORMAR DEVOLUÇÃO')
   db=st.session_state.db
   if db is None or db.empty:
    st.warning('A BASE DE MATERIAIS AINDA NÃO ESTÁ DISPONÍVEL.')
   else:
    _materials=db[['codigo','descricao']].drop_duplicates('codigo').copy()
    _materials=_materials.sort_values('codigo')
    _choices={
     f"{row.codigo} — {row.descricao}":(str(row.codigo),str(row.descricao))
     for _,row in _materials.iterrows()
    }
    with st.form('production_return_create_form',clear_on_submit=True):
     _picked=st.selectbox('MATERIAL',list(_choices.keys()),index=None,placeholder='Pesquise pelo código ou descrição')
     _q=st.number_input('QUANTIDADE',min_value=0.001,value=1.0,step=1.0,format='%.3f')
     _psy=st.text_input('PSY / PROJETO',placeholder='Informe a PSY vinculada à devolução')
     _obs=st.text_area('OBSERVAÇÃO DA DEVOLUÇÃO',height=90)
     _create=st.form_submit_button('REGISTRAR DEVOLUÇÃO',type='primary',use_container_width=True)
    if _create:
     if not _picked or not _psy.strip():
      st.error('MATERIAL, QUANTIDADE E PSY SÃO OBRIGATÓRIOS.')
     else:
      codigo,descricao=_choices[_picked]
      try:
       central_data.create_production_return(
        codigo,descricao,_q,_psy.strip().upper(),_obs.strip(),_session_auth_token(),
       )
       st.success('DEVOLUÇÃO REGISTRADA. AGUARDANDO RECEBIMENTO DO ALMOXARIFADO.')
       st.rerun()
      except Exception as exc:st.error(f'ERRO AO SOLICITAR DEVOLUÇÃO: {exc}')
  else:
   section_band('02 · CONFERÊNCIA','VALIDAR DEVOLUÇÕES NO ALMOXARIFADO')
   _waiting=[x for x in returns if x.get('status') in ('PENDENTE','RECEBIDO')]
   if not _waiting:
    st.success('NENHUMA DEVOLUÇÃO PENDENTE.')
   else:
    _options={
     f"{x.get('id','')[:8]} · {x.get('codigo','')} · {x.get('psy','')} · {x.get('status','')}":x
     for x in _waiting
    }
    with st.form('production_return_transition_form',clear_on_submit=True):
     _selected=st.selectbox('DEVOLUÇÃO PARA CONFERÊNCIA',list(_options.keys()))
     _item=_options[_selected]
     st.write(f"**{_item.get('descricao','')}** · Quantidade: **{fn(_item.get('quantidade') or 0)}** · PSY: **{_item.get('psy','')}**")
     if _item.get('observacao_producao'):
      st.caption('PRODUÇÃO: '+str(_item.get('observacao_producao')))
     _actions=['RECEBIDO','RECUSADO'] if _item.get('status')=='PENDENTE' else ['CONCLUIDO']
     _status=st.radio('VALIDAÇÃO',_actions,horizontal=True)
     _ref=st.text_input('REFERÊNCIA DO LANÇAMENTO PROTHEUS',placeholder='Obrigatório para concluir a devolução')
     _note=st.text_area('OBSERVAÇÃO DO ALMOXARIFADO',height=80)
     _save=st.form_submit_button('CONFIRMAR VALIDAÇÃO',type='primary',use_container_width=True)
    if _save:
     if _status=='CONCLUIDO' and not _ref.strip():st.error('INFORME A REFERÊNCIA DO LANÇAMENTO NO PROTHEUS.')
     elif _status=='RECUSADO' and not _note.strip():st.error('INFORME O MOTIVO DA RECUSA.')
     else:
      try:
       central_data.transition_production_return(
        str(_item['id']),_status,_note.strip(),_ref.strip(),_session_auth_token(),
       )
       st.success('VALIDAÇÃO REGISTRADA COM SUCESSO.')
       st.rerun()
      except Exception as exc:st.error(f'NÃO FOI POSSÍVEL VALIDAR A DEVOLUÇÃO: {exc}')
  topic_divider()
  section_band('03 · REGISTROS','RELATÓRIO DE DEVOLUÇÕES')
  if returns:
   _cols=['id','codigo','descricao','quantidade','psy','status','criado_por','criado_em',
          'recebido_por','recebido_em','concluido_por','concluido_em',
          'referencia_protheus','observacao_producao','observacao_almox']
   _report=pd.DataFrame(returns).reindex(columns=_cols).fillna('')
   _report=_report.rename(columns={
    'id':'ID','codigo':'CÓDIGO','descricao':'DESCRIÇÃO','quantidade':'QUANTIDADE',
    'psy':'PSY','status':'STATUS','criado_por':'SOLICITANTE','criado_em':'SOLICITADO EM',
    'recebido_por':'RECEBIDO POR','recebido_em':'RECEBIDO EM',
    'concluido_por':'CONCLUÍDO POR','concluido_em':'CONCLUÍDO EM',
    'referencia_protheus':'REF. PROTHEUS','observacao_producao':'OBS. PRODUÇÃO',
    'observacao_almox':'OBS. ALMOXARIFADO',
   })
   _report['QUANTIDADE']=pd.to_numeric(_report['QUANTIDADE'],errors='coerce').fillna(0)
   for _dc in ['SOLICITADO EM','RECEBIDO EM','CONCLUÍDO EM']:
    _report[_dc]=_report[_dc].map(lambda v:central_data.format_dt(v) if v else '')
   _setta_dataframe(_report,use_container_width=True,hide_index=True)
   st.download_button(
    'EXPORTAR REGISTROS EM EXCEL',excel_bytes(_report,'Retornos de Producao'),
    file_name='retornos_producao_setta.xlsx',
    mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    use_container_width=True,
   )
  else:st.info('AINDA NÃO EXISTEM DEVOLUÇÕES REGISTRADAS.')

# User administration
# Settings
elif active=='Configurações':
 if _session_profile()!='Gestor':
  section_band('01 · ACESSO','CONFIGURAÇÕES RESTRITAS','AUTENTIQUE UM USUÁRIO ADMINISTRADOR SETTA PARA ALTERAR CONFIGURAÇÕES.')
  if not _auth_user():
   with st.form('settings_auth_form',clear_on_submit=False):
    _settings_login=st.text_input('USUÁRIO OU E-MAIL',key='settings_auth_login')
    _settings_password=st.text_input('SENHA',type='password',key='settings_auth_password')
    _settings_submit=st.form_submit_button('IDENTIFICAR USUÁRIO',type='primary',use_container_width=True)
   if _settings_submit:
    try:
     if _authenticate_user(_settings_login,_settings_password):st.rerun()
     else:st.error('Usuário ou senha inválidos.')
    except Exception as exc:
     st.error(f'Falha ao autenticar: {exc}')
  else:
   st.warning('O usuário autenticado não possui perfil Gestor.')
  st.stop()
 tab_inv,tab_api=st.tabs(['INVENTÁRIO','ACOMPANHAMENTO DE API'])
 with tab_inv:
  section_band('01 · INVENTÁRIO','CONFIGURAÇÕES OPERACIONAIS')
  cfg['blind_default']=st.checkbox('CONTAGEM CEGA POR PADRÃO',cfg['blind_default'])
  if st.button('SALVAR CONFIGURAÇÕES',type='primary',use_container_width=True):
   if persist_cfg():st.success('CONFIGURAÇÕES SALVAS.')
   else:st.error('Não foi possível salvar as configurações. Refaça a autenticação e tente novamente.')
 with tab_api:
  render_api_monitor()

